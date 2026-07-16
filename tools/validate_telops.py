#!/usr/bin/env python3
"""コンテ表のテロップ列がテロップ記法v2に従っているか検証する。

使い方:
    python3 tools/validate_telops.py conte/第1回_コンテ表_v1.xlsx

方針: 疑わしい表記は「除去」せず「エラー」として人間に直させる。
パーサーが黙って本文を削ると、いつか画面に出すべき文字まで食われるため。
"""
import re
import sys

import openpyxl

COL_TELOP = 16  # テロップ(スタイル/内容)
COL_MOTION = 17  # テロップモーション

STYLE_RE = re.compile(r'^S[1-6](\(欧文\))?/')

# 本文への混入を疑う指示語(見つけたらエラー。本当に画面に出す文字なら記法を再考する)
FORBIDDEN = [
    '字幕全文', '併走', '画面主役', 'そのもの', '小テロップ',
    'のみ金茶', 'のみ大きく', 'に合わせ', 'スタイル', '同一', '呼応',
]


def validate(path: str) -> int:
    wb = openpyxl.load_workbook(path, read_only=True)
    if 'コンテ' not in wb.sheetnames:
        print(f'エラー: {path} に「コンテ」シートがない')
        return 1
    ws = wb['コンテ']
    problems = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        no = row[0]
        if not no:
            continue
        telop = row[COL_TELOP - 1] or ''
        motion = row[COL_MOTION - 1] or ''
        p_lines = [l for l in str(telop).split('\n') if l.strip()]
        for i, line in enumerate(p_lines, 1):
            tag = f'{no} テロップ{i}行目'
            if not STYLE_RE.match(line):
                problems.append(f'{tag}: 「Sn/本文」形式でない → {line}')
                continue
            body = line.split('/', 1)[1]
            if body == '全文':
                if not line.startswith('S1'):
                    problems.append(f'{tag}: 「全文」はS1のみ使用可 → {line}')
                continue
            if not body.strip():
                problems.append(f'{tag}: 本文が空 → {line}')
            # 「+20%」のような数値表記は合法。それ以外の+は連結の疑い
            if re.search(r'[+＋](?![0-9])', body):
                problems.append(f'{tag}: 「+」連結は禁止。1行1テロップに分ける → {line}')
            if re.search(r'[()（）]', body):
                problems.append(f'{tag}: 括弧内指示の疑い。指示はモーション列へ → {line}')
            for word in FORBIDDEN:
                if word in body:
                    problems.append(f'{tag}: 指示語「{word}」の混入疑い → {line}')
            # 本文中の別スタイル参照(例: 「〜」S5)
            if re.search(r'S[1-6]', body):
                problems.append(f'{tag}: 本文中にスタイル参照。行を分ける → {line}')
        # モーション列の行数対応(1行 or テロップと同数)
        q_lines = [l for l in str(motion).split('\n') if l.strip()]
        if q_lines and len(p_lines) > 1 and len(q_lines) not in (1, len(p_lines)):
            problems.append(
                f'{no}: モーション列の行数({len(q_lines)})がテロップ行数({len(p_lines)})と不一致'
            )

    if problems:
        print(f'✗ {path}: 違反 {len(problems)} 件')
        for p in problems:
            print('  -', p)
        return 1
    print(f'✓ {path}: 記法v2に適合(違反なし)')
    return 0


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(max(validate(p) for p in sys.argv[1:]))
