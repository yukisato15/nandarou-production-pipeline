"""MG Factoryの生成物に対する、AEを起動せず実行できる回帰テスト。"""

from __future__ import annotations

import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

from mg_factory.template_registry import IMPLEMENTED_TEMPLATES, validate_cut


ROOT = Path(__file__).resolve().parents[1]
MG_JSON = ROOT / "output/ep01_mg.json"
MG_JSX = ROOT / "mg_factory/generated/generate_mg_factory.jsx"
DESIGN_JSX = ROOT / "mg_factory/generated/generate_mg_factory_design_v2.jsx"
C13_MASTER_JSX = ROOT / "mg_factory/generated/generate_c13_master.jsx"
UI_HTML = ROOT / "mg_factory/generated/ui/C26_UI_AD_FEED_SCROLL.html"


class MgFactoryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.payload = json.loads(MG_JSON.read_text(encoding="utf-8"))
        cls.cuts = {cut["cut"]: cut for cut in cls.payload["cuts"]}

    def test_workbook_has_mg_columns_and_params(self) -> None:
        sheet = load_workbook(ROOT / "第1回_コンテ表_v1.xlsx", data_only=True, read_only=True)["コンテ"]
        headers = {str(cell.value): cell.column for cell in next(sheet.iter_rows(min_row=1, max_row=1)) if cell.value}
        self.assertIn("MGテンプレートID", headers)
        self.assertIn("MGパラメータJSON", headers)
        expected = {
            "C01", "C03", "C10", "C12", "C13", "C15", "C16", "C20", "C25",
            "C26", "C40", "C42", "C45", "C46", "C48", "C52", "C54", "C55",
        }
        for cut_no in expected:
            row = next(row for row in range(2, sheet.max_row + 1) if sheet.cell(row, headers["No"]).value == cut_no)
            params = json.loads(sheet.cell(row, headers["MGパラメータJSON"]).value)
            self.assertTrue(params)

    def test_json_contains_37_plans_and_18_implemented_cuts(self) -> None:
        self.assertEqual(len(self.payload["cuts"]), 37)
        expected = {
            "C01": "AE_CARD_STACK_THUMBNAIL",
            "C03": "AE_CARD_FOCUS_DROP",
            "C10": "AE_TITLE_EPISODE",
            "C12": "DATA_BAR_WORD_GROWTH",
            "C13": "DATA_FLOW_THREE_STEP",
            "C15": "DATA_RESEARCH_STAT_CARD",
            "C16": "DATA_LINE_COMPARE_TWO",
            "C20": "AE_BLACK_DECLARATION",
            "C25": "DATA_REPLACE_ARROW_DIAGRAM",
            "C26": "UI_AD_FEED_SCROLL",
            "C40": "AE_BLACK_DECLARATION",
            "C42": "AE_INTERSTITIAL_SIMPLE",
            "C45": "AE_DARK_STATEMENT",
            "C46": "AE_PRICE_TAG_MULTIPLY",
            "C48": "AE_LIST_ON_AMBIENT_PLATE",
            "C52": "AE_BRAND_LOGO_ASSEMBLY",
            "C54": "AE_NEXT_EPISODE_TITLE",
            "C55": "AE_END_CARD_CREDITS",
        }
        self.assertEqual({cut: self.cuts[cut]["template_id"] for cut in expected}, expected)
        self.assertEqual(sum(cut["template_id"] in IMPLEMENTED_TEMPLATES for cut in self.payload["cuts"]), 18)
        self.assertEqual((self.payload["project"]["width"], self.payload["project"]["height"]), (1920, 1080))

    def test_frame_math_is_consistent(self) -> None:
        for cut in self.payload["cuts"]:
            self.assertEqual(cut["duration_frames"], cut["out_frame"] - cut["in_frame"], cut["cut"])
            self.assertGreater(cut["duration_frames"], 0)

    def test_safety_flags_and_shared_master(self) -> None:
        self.assertTrue(self.cuts["C01"]["params"]["fictional"])
        self.assertTrue(self.cuts["C12"]["params"]["conceptual"])
        self.assertTrue(self.cuts["C16"]["params"]["conceptual"])
        self.assertIn("概念図", self.cuts["C16"]["params"]["note"])
        self.assertTrue(all(card["is_fictional"] for card in self.cuts["C26"]["params"]["cards"]))
        self.assertEqual(self.cuts["C20"]["params"]["reuse_key"], self.cuts["C40"]["params"]["reuse_key"])
        for tag in self.cuts["C46"]["params"]["tags"]:
            self.assertGreaterEqual(tag["x"], 0)
            self.assertLessEqual(tag["x"], 1)
            self.assertGreaterEqual(tag["y"], 0)
            self.assertLessEqual(tag["y"], 1)

    def test_jsx_is_bom_es3_and_embeds_implemented_cuts(self) -> None:
        raw = MG_JSX.read_bytes()
        self.assertEqual(raw[:3], b"\xef\xbb\xbf")
        jsx = raw.decode("utf-8-sig")
        self.assertIsNone(re.search(r"JSON\.parse\s*\(", jsx))
        self.assertIsNone(re.search(r"\b(?:const|let)\b|=>", jsx))
        for cut_no in (
            "C01", "C03", "C10", "C12", "C13", "C15", "C16", "C20", "C25",
            "C26", "C40", "C42", "C45", "C46", "C48", "C52", "C54", "C55",
        ):
            self.assertIn(f'"cut":"{cut_no}"', jsx)
        self.assertIn("cut.in_frame / FPS", jsx)
        self.assertIn('"reel_prefix":"EP01"', jsx)
        self.assertIn('"_MG_REVIEW_REEL"', jsx)
        self.assertIn('"_MG_TIMELINE_REEL"', jsx)
        self.assertIn('SHARED_COMPS[key]', jsx)
        self.assertIn('NANDAROU_MG_FACTORY_EP01', jsx)

    def test_registry_rejects_bad_nested_values(self) -> None:
        broken = json.loads(json.dumps(self.cuts["C46"], ensure_ascii=False))
        broken["params"]["tags"][0]["x"] = "left"
        warnings = validate_cut(broken)
        self.assertTrue(any("tags[1].x" in warning for warning in warnings))

    def test_partial_build_is_isolated_deterministic_and_check_is_read_only(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            first = temp / "first.jsx"
            second = temp / "second.jsx"
            command = [
                sys.executable, str(ROOT / "build_mg_factory.py"), str(MG_JSON),
                "--cuts", "C01,C12", "--output",
            ]
            subprocess.run(command + [str(first)], cwd=ROOT, check=True, capture_output=True, text=True)
            subprocess.run(command + [str(second)], cwd=ROOT, check=True, capture_output=True, text=True)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            jsx = first.read_text(encoding="utf-8-sig")
            self.assertIn("NANDAROU_MG_FACTORY_EP01_TEST_C01_C12", jsx)
            self.assertIn('"cut":"C01"', jsx)
            self.assertIn('"cut":"C12"', jsx)
            self.assertNotIn('"cut":"C10"', jsx)

            sentinel = temp / "sentinel.jsx"
            sentinel.write_text("KEEP", encoding="utf-8")
            subprocess.run(command + [str(sentinel), "--check"], cwd=ROOT, check=True, capture_output=True, text=True)
            self.assertEqual(sentinel.read_text(encoding="utf-8"), "KEEP")

    def test_committed_jsx_matches_current_builder(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            rebuilt = Path(directory) / "rebuilt.jsx"
            subprocess.run(
                [sys.executable, str(ROOT / "build_mg_factory.py"), str(MG_JSON), "--output", str(rebuilt)],
                cwd=ROOT,
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertEqual(MG_JSX.read_bytes(), rebuilt.read_bytes())

    def test_design_v2_is_isolated_es3_and_uses_safe_box_text(self) -> None:
        raw = DESIGN_JSX.read_bytes()
        self.assertEqual(raw[:3], b"\xef\xbb\xbf")
        jsx = raw.decode("utf-8-sig")
        self.assertIn('"cut":"C01"', jsx)
        self.assertIn('"cut":"C13"', jsx)
        self.assertIn('"cut":"C25"', jsx)
        self.assertNotIn('"cut":"C12"', jsx)
        self.assertIn('"NANDAROU_MG_FACTORY_" + PROJECT.episode + "_DESIGN_V2"', jsx)
        self.assertIn('property("ADBE Anchor Point").setValue([0, 0])', jsx)
        self.assertIsNone(re.search(r"\b(?:const|let)\b|=>", jsx))
        with tempfile.TemporaryDirectory() as directory:
            rebuilt = Path(directory) / "design_v2.jsx"
            subprocess.run(
                [sys.executable, str(ROOT / "build_mg_factory_design_v2.py"), str(MG_JSON), "--output", str(rebuilt)],
                cwd=ROOT,
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertEqual(raw, rebuilt.read_bytes())

    def test_c13_master_uses_eased_motion_and_four_second_board_hold(self) -> None:
        raw = C13_MASTER_JSX.read_bytes()
        self.assertEqual(raw[:3], b"\xef\xbb\xbf")
        jsx = raw.decode("utf-8-sig")
        self.assertIn('"duration_frames": 312', jsx)
        self.assertIn("new KeyframeEase(0, influence)", jsx)
        self.assertIn("setTemporalEaseAtKey", jsx)
        self.assertIn("function applyCameraRig(boardLayer, keys, fps, baseScale)", jsx)
        self.assertIn("function setTemporalEaseCompat(prop, keyIndex, inInfluence, outInfluence)", jsx)
        self.assertIn("for (dimensions = 1; dimensions <= 3; dimensions++)", jsx)
        self.assertIn("camera ease skipped:", jsx)
        self.assertIn('cam.name = "CAM_DRIFT"', jsx)
        self.assertIn("boardLayer.parent = cam", jsx)
        self.assertIn("pos.setValue([960, 540])", jsx)
        self.assertIn("{f: 76,  look: [780, 810],   zoom: 1.70}", jsx)
        self.assertIn("{f: 230, look: [1920, 1080], zoom: 1.00}", jsx)
        self.assertIn("ap.setSpatialTangentsAtKey(i, [0, 0], [0, 0])", jsx)
        self.assertIn("BG_MASTER_PAPER_CONTINUITY", jsx)
        self.assertIn("flash.inPoint = at(235)", jsx)
        self.assertIn("bg.inPoint = at(240)", jsx)
        self.assertIn("var emphasisGap = 44", jsx)
        self.assertIsNone(re.search(r"\b(?:const|let)\b|=>", jsx))

    def test_ui_is_standalone_and_deterministic(self) -> None:
        html = UI_HTML.read_text(encoding="utf-8")
        self.assertNotIn("http://", html)
        self.assertNotIn("https://", html)
        self.assertIn("window.MGPreview", html)
        self.assertIn("function renderFrame", html)
        self.assertIn("FICTIONAL UI", html)
        self.assertIn("MIRAI+", html)


if __name__ == "__main__":
    unittest.main()
