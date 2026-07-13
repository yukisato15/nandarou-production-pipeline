# -*- coding: utf-8 -*-
"""第2回コンテ表を第1回のワークブックを雛形に生成する。
方針: 自作MG(AE新規構築)ゼロ。AEマスター複製 / Remotion生成 / HTML・UIモック /
Canva+AI静止画 / Envato / 実撮影 / 画面収録 のみで構成する。
"""
import openpyxl, math, sys

SRC = "第1回_コンテ表_v1.xlsx"
DST = "第2回_コンテ表_v1.xlsx"

HEADERS = ['No','パート','TC目安','ナレーション(全文)','読み方の注意','映像内容(画)','素材ソース',
 'Envato検索ワード(映像)','Runway生成プロンプト','Runwayネガティブ','BGM','SE','SE Envato検索',
 'エフェクト(+Envato検索)','次カットへのトランジション(+Envato検索)','テロップ(スタイル/内容)',
 'テロップモーション','テロップFX/TRN Envato検索','MGテンプレートID','MG自動化内容','MG必要入力',
 'MG補助ツール・素材','MG手動確認・注意','MGパラメータJSON']

CPS = 5.5

def cut(no, part, nar, visual, source, *, extra=1.2, fixed=None, yomi='—', envato='—',
        runway='—', runway_ng='—', bgm='M1', se='—', se_q='—', fx='—', trn='—',
        telop='—', tmotion='—', tfx='—', mgid='—', mgauto='—', mgin='—', mgtool='—',
        mgcheck='—', mgjson='—'):
    return dict(no=no, part=part, nar=nar, visual=visual, source=source, extra=extra,
        fixed=fixed, yomi=yomi, envato=envato, runway=runway, runway_ng=runway_ng,
        bgm=bgm, se=se, se_q=se_q, fx=fx, trn=trn, telop=telop, tmotion=tmotion,
        tfx=tfx, mgid=mgid, mgauto=mgauto, mgin=mgin, mgtool=mgtool, mgcheck=mgcheck,
        mgjson=mgjson)

C = []

# ============================================================ P0 フック
C.append(cut('C01','P0',
 '「は?」。\n「ありえない」。\n「拡散希望」。',
 '黒背景に、炎上している架空のSNSポスト(引用リポスト・怒りのコメント)が通知音と同時に次々スタック。1件ごとに画面が埋まっていく(第1回C01と同一機構)',
 'AEマスター複製', extra=6, bgm='M0',
 yomi='読まない選択肢もあり(テロップ+SEのみ)。読む場合は感情ゼロの棒読みで"引用"だと分からせる',
 se='通知音連打/ポスト着弾のドン(低め)', se_q='notification pop UI\nsoft impact hit low',
 fx='ポストにVHS風ノイズ薄く=「VHS overlay texture」',
 trn='カット尻で全ポスト微ズーム寄り→C02へハードカット',
 telop='S6/ポスト文言そのもの(は? 等)', tmotion='1件ごとにスケール105→100%のドン付き出現 [maxchar:12]',
 mgid='AE_CARD_STACK_THUMBNAIL',
 mgauto='第1回のカードスタックマスターを複製し、画像をSNSポスト風PNGへ差し替え。出現delay・ドン・最終微ズームは流用。',
 mgin='ポスト画像3〜5枚、表示順、出現間隔',
 mgtool='ポスト画像はCanvaで架空アカウント名・架空アイコンで作成。AI生成アイコン(日本人・顔認識不能レベル)。',
 mgcheck='実在アカウント・実在事件を連想させる文言になっていないか確認。',
 mgjson='{"fictional":true,"reuse_master":"EP01_C01","cards":5}'))

C.append(cut('C02','P0',
 '——SNSを閉じたあとに、妙に疲れていること、ないでしょうか。',
 '夜、ベッドサイド。スマホを伏せて置く手。画面の光が消えて、部屋が暗くなる',
 'Envato', bgm='M0',
 envato='putting smartphone down night bed\nphone screen light face dark room',
 trn='画面の光が消える暗転をそのままカット点に使う',
 telop='S1/字幕全文'))

C.append(cut('C03','P0',
 '白状すると、この動画のタイトルも、あなたがちょっとムッとするように、作っています。',
 'この動画自身のサムネが画面中央に落ちてきて静止。周囲の炎上ポストは暗く沈む(第1回C03と同一機構)',
 'AEマスター複製', bgm='M0',
 se='着地のドン(単発)', se_q='soft impact hit low',
 telop='S1/字幕全文',
 mgid='AE_CARD_STACK_THUMBNAIL',
 mgauto='第1回C03マスター複製。中央カード画像を本編サムネへ差し替え。',
 mgin='本編サムネ画像1枚', mgtool='サムネはCanva/Photoshopで制作(偽サムネと同工程)。',
 mgcheck='サムネコピーと本編内容の裏付け一致を確認。'))

C.append(cut('C04','P0',
 '前回は、「不安」の話をしました。不安を煽ると再生数が伸びて、広告のお金に変わる——という仕組みの話です。',
 '第1回C13の分解図ボード(目→再生→広告費)を静止画で再掲。ゆっくり寄り(前回の復習であることが一目で分かる画)',
 'AEマスター複製',
 telop='S3/前回:「不安」の仕組み',
 mgid='AE_KAITAI_BOARD_C13',
 mgauto='第1回C13マスターをそのまま複製し、全景静止+微ズームのみに変更。',
 mgin='なし(完成コンポ流用)', mgcheck='第1回の画とトーンが揃っているか確認。'))

C.append(cut('C05','P0',
 '今回は、その続きです。不安よりも、もっとよく燃える燃料——「怒り」の話をします。',
 'ボードの紙が朱色にじわっと染まっていき、暗転。コメント欄UIモックが赤く明滅して立ち上がる',
 'HTML/UI生成',
 fx='朱のインクにじみ=「ink bleed transition」',
 trn='紙フラッシュ→UIモックへ',
 telop='S2/★不安の、その先へ。', tmotion='★文:中央フェード',
 mgid='HTML_COMMENT_FEED',
 mgauto='架空コメント欄UIをHTMLで生成し画面収録。コメント文言はJSONから流し込み。',
 mgin='コメント文言リスト、明滅タイミング', mgtool='HTML/CSSモック+画面収録。実在サービスのUIを直接模倣しない。',
 mgcheck='実在サービスのロゴ・配色を避ける。'))

C.append(cut('C06','P0',
 'テレビにも、人を怒らせる技術は、ありました。ワイドショーです。コメンテーターが眉をひそめるたび、視聴率は動きました。',
 '古いブラウン管テレビ(レトロ質感)。ワイドショー風の抽象画面(実番組は映さない)。スタジオの照明',
 'Envato',
 envato='retro CRT television living room\nold tv static vintage broadcast',
 telop='S1/字幕全文'))

C.append(cut('C07','P0',
 'でも、テレビの怒りは、翌朝には冷めました。いまの怒りは、朝も昼も夜中も、ポケットの中で燃えつづけています。',
 '夜の寝室。枕元のスマホに通知の光が点き続ける。時計が2時、3時と進む',
 'Envato実写+AE合成',
 envato='smartphone notification glow dark bedroom night\nclock time lapse night',
 fx='通知光の明滅を合成(既存LOOK_GLOBALの光素材)',
 telop='S1/字幕全文'))

C.append(cut('C08','P0',
 'なぜ、こんなに、みんな怒っているんだろう。',
 '暗がりで、顔に赤いコメント欄の光が反射している抽象イメージ(顔は判別不能レベル)',
 'AI生成(Runway)',
 runway='a person in a dark room, face barely visible, red light from a screen reflecting on their silhouette, abstract, cinematic, shallow depth of field',
 runway_ng='recognizable face, text, logo, bright colors',
 telop='S1/字幕全文'))

C.append(cut('C09','P0',
 '★調べてみると——あなたの怒りには、値段が、ついていました。',
 '机の上のノート、付箋、資料、コーヒー。調べものの手元(Yukiの実物)。付箋の1枚に小さく「¥」',
 '実撮影(推奨)',
 yomi='「値段が、ついていました」は張らずに、間で落とす',
 telop='S2/★あなたの怒りには、値段がついている。', tmotion='★文:中央。「値段」のみ金茶1.2倍'))

C.append(cut('C10','P0',
 '第2回。——あなたの怒りは、誰の収益になっているのか。ここから、始めます。',
 'タイトルカード:墨背景に「怒りは、燃料になる」+小さく『なんだろう』解体 #02',
 'AEマスター複製', extra=4,
 se='紙フラッシュ', se_q='paper whoosh transition',
 telop='S4/怒りは、燃料になる(タイトル内)',
 mgid='AE_TITLE_CARD',
 mgauto='第1回タイトルカードマスター複製。題字と回数のみ差し替え。',
 mgin='題字、回数', mgcheck='題字の改行位置を確認。'))

# ============================================================ P1 現象・仕組み
C.append(cut('C11','P1',
 'まず、結論から言います。\n★怒りは、感情ではありません。燃料です。',
 '暗い部屋でスマホをスクロールする指。画面の光だけが顔の輪郭を照らす(第1回C11と対の画)',
 'Envato',
 envato='scrolling smartphone dark room face glow\nthumb scrolling social media close up',
 telop='S2/★怒りは、感情ではなく燃料です。', tmotion='★文:中央。「燃料」のみ金茶+1.2倍'))

C.append(cut('C12','P1',
 'ある研究チームが、SNSの投稿を50万件、調べました。「許せない」「卑怯だ」——道徳と感情の混ざった言葉が1語増えるごとに、拡散は約2割、増えていました。',
 '紙資料風の棒グラフ:語数0〜3語に対して拡散率の棒が伸びる。数字カウントアップ。出典表記を規定位置に常設',
 'Remotion生成', extra=4,
 telop='S5/+20%(数字)/1語ごとに(和文)', tmotion='数字カウントアップ、棒はインク質感で伸長',
 mgid='RMT_BAR_MORAL_WORDS',
 mgauto='RemotionのデータMGキット(棒グラフ)にJSONを流し込み、透過ProResで書き出し→AEでLOOK_GLOBAL仕上げ。',
 mgin='棒の値配列、ラベル、出典文字列、尺',
 mgtool='remotion_mg/ を拡張(Codex作成のC13AttentionFlowと同一プロジェクト)。',
 mgcheck='「約2割」の留保表現が画面とナレで一致しているか。出典は概要欄と一致させる。',
 mgjson='{"engine":"remotion","comp":"BarMoralWords","source":"Brady et al. 2017 (研究の要約・留保形)"}'))

C.append(cut('C13','P1',
 '怒った人は、シェアします。シェアされた投稿は、また誰かを怒らせます。その間ずっと、画面には広告が出続けています。★回っているのは、議論ではなくて、広告なんです。',
 '紙のボード分解図:「怒り」→「シェア」→「広告費」。左=朱に染まった吹き出し/中=リポストUI風カード/右=¥紙片。金茶の線で接続(第1回C13マスターの複製・素材差し替え)',
 'AEマスター複製', extra=4,
 se='紙めくり/スタンプ', se_q='paper flip\nstamp thud',
 telop='S2/★回っているのは、広告です。', tmotion='★文:全景の引きと同期して中央',
 mgid='AE_KAITAI_BOARD_C13',
 mgauto='C13マスター複製。3ノードの画像・ラベル文言を差し替え。アンカーカメラリグは流用。',
 mgin='ノード画像3点、ラベル3語、接続線タイミング',
 mgtool='吹き出し・リポストカードはCanva+HTMLモックの静止画。¥紙片は第1回素材を流用。',
 mgcheck='因果を断定しすぎない文言か確認。'))

C.append(cut('C14','P1',
 'ひとつ、作る側の話をします。映像は、切る場所で意味が変わります。同じインタビューでも、前後を落とせば、真逆に聞こえる。ぼくらはこれを「切り取り」と呼んで、やらないよう教わりました。——いまは、「伸びる技術」と呼ばれています。',
 '編集ソフトのタイムライン実画面。同一インタビュー素材(自撮り・架空内容)の切り取りBefore/Afterで印象が反転する',
 '画面収録', extra=4,
 yomi='「伸びる技術」は皮肉を込めず、平温で',
 telop='S1/字幕全文(Before/Afterに同期)'))

C.append(cut('C15','P1',
 '気のせいでは、ありません。2021年、大手SNSの内部文書が、報道機関に渡りました。そこには——「怒り」の絵文字を、「いいね」の5倍、重く扱う、と書かれていました。',
 '紙資料風の数字MG:😡(怒り)=👍(いいね)×5 の等式が、タイプライター的に組み上がる。「内部文書・報道による」の留保を画面内に常設',
 'Remotion生成', extra=4,
 telop='S5/×5(数字)/いいねの5倍(和文)', tmotion='等式が順に出現、×5のみ朱',
 mgid='RMT_EQUATION_WEIGHT',
 mgauto='Remotion数式・カウンター部品にJSON流し込み。透過書き出し→AE仕上げ。',
 mgin='左右の項、係数、留保注記', mgtool='remotion_mg/ に等式コンポを追加。',
 mgcheck='「報道によれば」の留保形を必ず画面に残す(断定しない)。',
 mgjson='{"engine":"remotion","comp":"EquationWeight","value":5,"caveat":"2021年・米紙報道の内部文書に基づく"}'))

C.append(cut('C16','P1',
 '重く扱われた投稿は、より多くの人のフィードに載ります。つまり——いちばん怒らせた投稿が、いちばん多く、配られていた。',
 '紙資料風の折れ線2本:「怒り投稿の到達」(朱)が「通常投稿」(墨)を引き離していく(第1回C16と対のグラフ)',
 'Remotion生成',
 telop='S3/いちばん怒らせた投稿が、いちばん配られる',
 mgid='RMT_LINE_COMPARE',
 mgauto='Remotion折れ線部品(2系列)にJSON流し込み。第1回C16と同型・配色規定準拠。',
 mgin='2系列の値配列、ラベル、尺', mgtool='remotion_mg/。',
 mgcheck='概念図であることを注記(値は演出用)。',
 mgjson='{"engine":"remotion","comp":"LineCompare","conceptual":true}'))

C.append(cut('C17','P1',
 '楽しい話は、「へえ」で終わります。でも「許せない」は、指が止まらない。つい引用して、ひとこと、書きたくなる。',
 '引用ポストが枝分かれして増殖していくUIモック(第1回C17の複製・文言差し替え)',
 'AEマスター複製',
 se='ポップ音の連鎖', se_q='ui pop bubble chain',
 telop='S1/字幕全文',
 mgid='HTML_QUOTE_CASCADE',
 mgauto='第1回C17のHTMLモックを複製し、吹き出し文言をJSONから差し替え。',
 mgin='吹き出し文言リスト', mgcheck='文言が実在の炎上を特定しないこと。'))

C.append(cut('C18','P1',
 'しかも——怒った投稿に「いいね」がつくと、人は、もっと怒った投稿をするようになります。★怒りっぽいのは、性格ではありません。学習です。',
 '黒背景+テロップのみ(M6宣言)',
 'AEマスター複製', bgm='M0',
 yomi='「学習です」で切る。間を編集で作る',
 telop='S2/★怒りっぽさは、性格ではなく学習です。', tmotion='1文字2Fずらし表示(M6規定)',
 mgid='AE_M6_STATEMENT',
 mgauto='M6宣言マスター複製。文言差し替えのみ。',
 mgin='★文、強調語', mgcheck='研究の言及は概要欄出典と一致させる。'))

C.append(cut('C19','P1',
 'そして、その学習を、収益に変えている人たちがいます。',
 '架空の収益ダッシュボードUI(インプレッション数、収益額のカウンタが回る)モック',
 'HTML/UI生成',
 telop='S1/字幕全文',
 mgid='HTML_REVENUE_DASH',
 mgauto='架空ダッシュボードをHTMLで生成し画面収録。数値はJSONから。',
 mgin='表示数値、通貨表記', mgcheck='実在サービスのUIを直接模倣しない。'))

C.append(cut('C20','P1',
 '今日の後半では、そこまで行きます。ひとつだけ、先に言っておくと。\n★わざと、嫌われている人が、います。',
 '完全な黒画面。テロップのみ(第1回C20と同一運用)',
 'AEマスター複製', bgm='M0',
 telop='S2/★わざと嫌われている人が、います。', tmotion='M6規定',
 mgid='AE_M6_STATEMENT', mgauto='M6宣言マスター複製。', mgin='★文'))

C.append(cut('C21','P1',
 'その話は、あとで。まずは、130年前に遡ります。怒りを売る商売は、SNSの発明では、ないからです。',
 'コメント欄UIが白黒に反転→セピアの新聞紙面へ。時代を遡るトランジション(第1回C21の複製)',
 'AEマスター複製',
 se='フィルムのコマ飛び', se_q='film projector rewind',
 trn='セピア反転トランジション(第1回素材流用)',
 telop='S1/字幕全文'))

# ============================================================ P2 歴史
C.append(cut('C22','P2',
 '1890年代のニューヨーク。2つの新聞社が、部数競争をしていました。ピュリツァーと、ハースト。——いま報道の賞に名前が残る人たちが、当時やっていたのは。',
 '1890年代ニューヨーク。新聞街、輪転機、新聞売りの少年(セピア・フィルム質感)',
 'AI生成(Runway)', bgm='M2',
 runway='1890s New York newspaper district, printing presses, newsboys on the street, sepia tone, old film grain, cinematic wide shot',
 runway_ng='modern buildings, cars, color, readable text, real logos',
 telop='S3/1890年代・ニューヨーク'))

C.append(cut('C23','P2',
 'でかい見出しで、怒りと恐怖を売ることでした。事件は、「ある」ことより、「燃える」ことが大事。のちに「イエロージャーナリズム」と呼ばれる手法です。',
 '架空の扇情紙面(英字・巨大見出し・実在紙面は使わない)が次々に刷り上がる',
 'Canva+AI静止画', bgm='M2',
 fx='紙の質感・輪転機の動きは2.5Dパララックス',
 telop='S3/イエロージャーナリズム',
 mgid='AI_25D_COLLAGE',
 mgauto='架空紙面(Canvaで組版)+AI背景をレイヤー分割し、2.5Dプリセットで動かす(第1回M4方式)。',
 mgin='紙面画像2〜3枚、パララックス量', mgtool='紙面はCanvaで英字架空見出し。書体は当時風セリフ。',
 mgcheck='実在紙名・実在事件見出しの再現になっていないか確認。'))

C.append(cut('C24','P2',
 '戦争の写真が無ければ、絵描きに描かせて載せる。「君は絵を用意しろ、戦争は私が用意する」。ハーストがそう言った、という逸話まであります。実話かは、怪しい。でも、そんな逸話が生まれるほどの、商売だったということです。',
 '電信室と画家の手元(抽象・セピア)。描かれた戦場の絵が紙面にはめ込まれる',
 'AI生成(Runway)', bgm='M2',
 yomi='逸話部分は引用の棒読み。「実話かどうかは、怪しい」は平温で',
 runway='1890s telegraph office, an illustrator\'s hands drawing a war scene, sepia, film grain, abstract close-up, no faces',
 runway_ng='recognizable face, modern objects, color, text',
 telop='S1/字幕全文'))

C.append(cut('C25','P2',
 '★新聞が売っていたのは、事実ではなくて、興奮でした。',
 '紙ボードの図解:「事実」に朱の筆致で打ち消し→金茶の手描き矢印→「興奮」(第1回C25マスターの複製・語差し替え)',
 'AEマスター複製', bgm='M2',
 se='筆致のストローク/ピン留め', se_q='brush stroke swoosh\npin push',
 telop='S2/★売っていたのは「興奮」です。', tmotion='★文:図解と同期して中央下 [pos:bottom]',
 mgid='AE_KAITAI_REPLACE',
 mgauto='第1回C25(打ち消し→矢印)マスター複製。左右の語のみ差し替え。',
 mgin='左語、右語', mgcheck='左右の意味関係と改行を確認。'))

C.append(cut('C26','P2',
 '——昔話に、聞こえるでしょうか。今日の、トレンド欄と見比べてみてください。',
 '一瞬で現代へ:トレンド欄風UIモック(架空ワード)。彩度が戻る(第1回C26と同運用)',
 'HTML/UI生成',
 trn='セピア→カラーの鋭いマッチカット',
 telop='S1/字幕全文',
 mgid='HTML_TREND_MOCK',
 mgauto='架空トレンド欄をHTMLで生成し画面収録。ワードはJSONから。',
 mgin='トレンドワード8件', mgcheck='実在の炎上ワード・実名を使わない。'))

C.append(cut('C27','P2',
 '1930年代は、ラジオでした。アメリカのある神父は、憎しみのこもった説教で、数千万の聴取者を集めた、とされています。名指しされた「敵」への攻撃は、放送のたびに増えました。',
 '1930年代の居間。家族が囲む古いラジオ(シルエット・逆光・顔なし)。ラジオの目盛りの光',
 'AI生成(Runway)', bgm='M2',
 runway='1930s living room, family silhouettes around an old radio, warm dim light, sepia tone, film grain, no faces visible',
 runway_ng='recognizable face, modern objects, text, logo',
 telop='S3/1930年代・ラジオ'))

C.append(cut('C28','P2',
 'テレビは、これを「討論」の形にしました。対立が激しいほど、数字が取れる。だから、いちばん極端な人を、両端に座らせる。',
 'テレビスタジオの抽象(対峙する2つの椅子、照明、カメラのシルエット。実番組は映さない)',
 'AI生成(Runway)', bgm='M2',
 runway='empty tv studio, two chairs facing each other, dramatic spotlights, silhouettes of cameras, dark, cinematic, abstract',
 runway_ng='people, faces, network logos, text',
 telop='S1/字幕全文'))

C.append(cut('C29','P2',
 'インターネットの最初の発明は、匿名掲示板でした。次の発明が——「まとめ」です。ケンカのログを読みやすく整えて、広告を貼る。★他人の怒りが、初めて、そのまま商品になりました。',
 '匿名掲示板風→まとめ記事風のUIモック(架空)。ケンカのログに広告枠が挟まっていく',
 'HTML/UI生成', bgm='M2',
 telop='S2/★他人の怒りが、商品になった。', tmotion='★文:UIの上に中央 [pos:center]',
 mgid='HTML_MATOME_MOCK',
 mgauto='掲示板→まとめ記事のUIをHTMLで生成し画面収録。本文・広告枠はJSONから。',
 mgin='ログ文言、広告枠位置', mgcheck='実在サイトのデザイン・実在スレの再現をしない。'))

C.append(cut('C30','P2',
 'ここまでは、まだ、人間が怒りを選んでいました。編集長が。プロデューサーが。管理人が。',
 '紙面・スタジオ・掲示板画面の3連モンタージュ(既出素材の再利用)。それぞれに「選ぶ手」のインサート',
 'Envato実写+AE合成', bgm='M2',
 envato='editor desk newspaper vintage\nhand choosing photo lightbox',
 telop='S1/字幕全文'))

C.append(cut('C31','P2',
 'いまは、違います。★怒る相手を、機械が、あなた専用に選んでくるんです。',
 '暗闇に無数のスマホの光。1台だけがこちらを向いている(第1回C33の複製)',
 'AEマスター複製', bgm='M2',
 telop='S2/★怒る相手は、機械が選んでいる。', tmotion='M6規定に準ずる',
 mgid='AE_PHONE_FIELD',
 mgauto='第1回C33の合成マスター複製。手前1台の画面内容のみ差し替え。',
 mgin='手前スマホの画面画像1枚'))

C.append(cut('C32','P2',
 'あなたが立ち止まった投稿。最後まで読んだ口論。開いたコメント欄。ぜんぶ記録されて、「この人は、これに怒る」というプロフィールになります。次の怒りは、それを見て届きます。',
 '行動ログ→プロフィール→配信、の3画面が順に繋がるUIモック(架空)。第1回C34の方式を流用',
 'HTML/UI生成', bgm='M2', extra=3,
 telop='S1/字幕全文',
 mgid='HTML_PROFILE_FLOW',
 mgauto='3画面のUIモックをHTMLで生成し画面収録。項目はJSONから。',
 mgin='ログ項目、プロフィール項目', mgcheck='実在サービス名を出さない。'))

C.append(cut('C33','P2',
 '日本でも、起きています。ある人物への中傷動画が、何年も収益を上げつづけた事案。デマがコピーされ、削除されても、別のアカウントが上げなおす。——本人が、否定しても、です。',
 '国会議事堂や裁判所の引きの実写(昼・無人寄り)。資料紙片のインサート(固有名なし)',
 'Envato', bgm='M2', extra=3,
 yomi='固有名は出さない。三点留保(本人否定/報道側訂正/未確定)を維持した読みに',
 envato='japan national diet building wide shot\ncourthouse exterior calm',
 telop='S1/字幕全文',
 mgcheck='法務:固有名なし・特定可能な特徴なし・三点留保。'))

C.append(cut('C34','P2',
 'なぜ、止まらないのか。上げなおす側に、悪意すら、要らないからです。再生されれば、収益になる。それだけで、十分に、回ってしまう。',
 '同じサムネ(架空)が複製されて増えていくグリッドのUIモック',
 'HTML/UI生成', bgm='M2',
 telop='S1/字幕全文',
 mgid='HTML_CLONE_GRID',
 mgauto='サムネ複製グリッドをHTMLで生成し画面収録。',
 mgin='サムネ画像1枚、複製数・間隔'))

C.append(cut('C35','P2',
 '★怒りは、いちばん安い燃料なんです。取材も、事実確認も、要らないから。',
 '黒背景+テロップのみ(M6宣言)',
 'AEマスター複製', bgm='M0',
 telop='S2/★怒りは、いちばん安い燃料。', tmotion='M6規定',
 mgid='AE_M6_STATEMENT', mgauto='M6宣言マスター複製。', mgin='★文'))

# ============================================================ P3 構造の核心
C.append(cut('C36','P3',
 'さて——さっきの予告を、回収します。',
 '静かな日常実写(昼の部屋、窓、湯気)で一拍おく',
 'Envato', bgm='M1',
 envato='quiet room window light steam coffee',
 telop='S1/字幕全文'))

C.append(cut('C37','P3',
 '★わざと、嫌われている人が、います。',
 'C20と完全に同じ黒画面+テロップ(同じ画で回収する)',
 'AEマスター複製', bgm='M0',
 telop='S2/★わざと嫌われている人が、います。', tmotion='M6規定(C20と同一)',
 mgid='AE_M6_STATEMENT', mgauto='C20のコンポをそのまま再使用。'))

C.append(cut('C38','P3',
 'わざと下手な料理を作る。わざと間違ったことを言う。わざと、怒らせる。コメント欄が「は?」で埋まるほど、動画は伸びます。憎まれることが、そのまま再生数になるからです。',
 '架空の「わざと系」動画のUIモック(料理動画風・顔なし手元)+荒れるコメント欄。すべて架空に再構成',
 'Canva+AI静止画', bgm='M1', extra=3,
 telop='S1/字幕全文',
 mgid='HTML_RAGE_VIDEO_MOCK',
 mgauto='架空動画ページをHTML+Canva素材で構成し画面収録。コメントはJSONから。',
 mgin='動画サムネ(架空)、コメント文言',
 mgtool='料理の手元はAI静止画または自前撮影。実在の投稿者を連想させる要素は排除。',
 mgcheck='法務:実在クリエイターの特定に繋がる特徴(口癖・構図・料理名の組合せ)を避ける。'))

C.append(cut('C39','P3',
 '最近は、名前もついています。「レイジベイト」——怒りの、撒き餌です。あなたの「ひとこと言ってやりたい」が、彼らの設計の、最後のピースです。',
 '紙ボード分解図:釣り針に吹き出し(「は?」)が食いつく図。金茶の線で「撒き餌→食いつき→収益」(C13ボードのレイアウトバリアント)',
 'AEマスター複製', bgm='M1', extra=3,
 se='ピン留め/紙めくり', se_q='pin push\npaper flip',
 telop='S3/レイジベイト=怒りの撒き餌',
 mgid='AE_KAITAI_BOARD_C13',
 mgauto='C13ボードマスター複製。ノード画像とラベルを差し替え(釣り針・吹き出し・¥)。',
 mgin='ノード画像3点、ラベル', mgtool='釣り針・吹き出しはAI静止画+切り抜き。',
 mgcheck='構図がC13の使い回しに見えないよう、ノード配置は左右反転などで変化をつける。'))

C.append(cut('C40','P3',
 'しかも、いまのSNSは、表示回数に応じて、投稿者にお金を配ります。怒りのリプライも、批判の引用も、表示は表示。ぜんぶ——★燃やした本人の、売上です。',
 '紙資料風カウンターMG:表示回数が回り、一定ごとに¥へ変換されていく',
 'Remotion生成', extra=3,
 telop='S5/¥(数字)/表示回数→収益(和文)', tmotion='カウントアップ→¥変換で朱が一瞬差す',
 mgid='RMT_COUNTER_FLOW',
 mgauto='Remotionカウンター部品(C13AttentionFlowの再生カウンターを流用)にJSON流し込み。',
 mgin='初期値、増加率、変換レート表示', mgtool='remotion_mg/ 既存コンポの改修。',
 mgcheck='具体的な収益単価を断定しない(制度は変わるため)。',
 mgjson='{"engine":"remotion","comp":"CounterFlow","reuse":"C13AttentionFlow"}'))

C.append(cut('C41','P3',
 'あなたは、抗議したつもりでした。でも、構造の側から見ると——あなたは、無料で、広告を運んだんです。',
 '吹き出し(リプライ)のひとつひとつに値札タグが貼られていく(第1回C46値札マスターの応用)',
 'AEマスター複製', bgm='M1',
 se='タグを貼る音', se_q='price tag snap',
 telop='S1/字幕全文',
 mgid='AE_PRICE_TAG_OVERLAY',
 mgauto='第1回C46の値札増殖マスター複製。貼り付け先を吹き出し画像に差し替え。',
 mgin='吹き出し画像群、タグ文言'))

C.append(cut('C42','P3',
 '★怒っているあいだ、いちばん静かに得をしている人を、探してください。',
 '黒背景+テロップのみ(M6宣言)',
 'AEマスター複製', bgm='M0',
 yomi='張らない。静かに置く',
 telop='S2/★いちばん静かに得をしている人は、誰か。', tmotion='M6規定',
 mgid='AE_M6_STATEMENT', mgauto='M6宣言マスター複製。', mgin='★文'))

# ============================================================ P4 まとめ・防衛
C.append(cut('C43','P4',
 '今日、覚えて帰ってほしいことは、三つです。',
 '朝の街、窓から差す光、湯気——ニュートラルで静かな生活の画(第1回C47と同トーン)',
 'Envato', bgm='M3',
 envato='morning city calm window light\nsteam kettle morning kitchen',
 telop='S1/字幕全文'))

C.append(cut('C44','P4',
 '★ひとつ。怒りは、いちばん速く広がって、いちばん安くつく燃料です。\n★ふたつ。あなたが怒る相手は、機械が、あなた専用に選んでいます。\n★みっつ。あなたの怒りのリプライは、燃やした本人の売上になります。',
 '静かな背景の上に3点がテロップで順に積まれる(第1回C48マスター複製)',
 'AEマスター複製', bgm='M3', extra=4,
 telop='S2/★3点(順次)', tmotion='順に積む。各行の強調語のみ金茶',
 mgid='AE_THREE_POINTS',
 mgauto='第1回C48の3点積みマスター複製。文言差し替え。',
 mgin='3文、強調語3つ'))

C.append(cut('C45','P4',
 'ぼくも、調べるまでは、見るたびに腹を立てていました。いまは、怒る前に、一拍だけ置きます。——「これで、誰が儲かる?」。それだけです。',
 'スマホを机に置き、ノートに一行書く手元(Yukiの実物)。書かれるのは「誰が儲かる?」',
 '実撮影(推奨)', bgm='M3', extra=3,
 yomi='「それだけです」は軽く。説教にしない',
 telop='S1/字幕全文'))

C.append(cut('C46','P4',
 '怒っては いけない、という話ではありません。★怒りの持って行き先を、あなたが決める、という話です。',
 '窓から差す光がゆっくり動く(第1回C50と同トーン)',
 'Envato', bgm='M3',
 envato='window light moving time lapse room calm',
 telop='S2/★怒りの行き先は、自分で決める。', tmotion='★文:中央フェード'))

C.append(cut('C47','P4',
 '本当に大事な怒りは、リプライ欄ではないところで、使ってください。署名かもしれない。投票かもしれない。目の前の誰かへの、一本の電話かもしれません。',
 '夕方の街。歩く人々、店の灯り、家路(第1回C51と同トーン。感傷的にしすぎない普通の夕方)',
 'Envato', bgm='M3', extra=3,
 envato='evening street people walking home japan\nshop lights dusk town',
 telop='S1/字幕全文'))

# ============================================================ P5 チャンネル・次回
C.append(cut('C48','P5',
 'このチャンネル——『なんだろう』解体では、こういう「なんだろう」を、ひとつずつ、構造から解体していきます。教えるんじゃなくて、一緒に、調べていく。',
 '『なんだろう』解体キービジュアル(第1回C52マスター複製)',
 'AEマスター複製', bgm='M3',
 telop='S4/『なんだろう』解体',
 mgid='AE_CHANNEL_KV', mgauto='第1回C52マスターをそのまま再使用。'))

C.append(cut('C49','P5',
 '今回も、台本と出典は、概要欄からぜんぶ読めます。どこまでが研究や報道で、どこからがぼくの解釈か。間違っていたら、次の動画の頭で、訂正します。',
 'noteの記事画面(第2回の出典リスト)をスクロールする画面収録',
 '画面収録', bgm='M3',
 telop='S1/字幕全文',
 mgcheck='出典リストは公開前に本編の言及と突合する。'))

C.append(cut('C50','P5',
 '次回は——★「『おすすめ』は、なぜ、あなたの好みを知っているのか」。\n見たいものが、先回りして届く、あの気持ち悪さ。仕組みから、解体します。\nそれでは、また次回。',
 '黒地に次回タイトルのテロップ+粒子の揺らぎ(第1回C54マスター複製)【次回テーマは仮。確定後に差し替え】',
 'AEマスター複製', bgm='M3',
 telop='S2/★次回:「おすすめ」はなぜ当たるのか(仮)', tmotion='M6規定+粒子',
 mgid='AE_NEXT_PREVIEW',
 mgauto='次回予告マスター複製。次回番号・題名・問いのみ差し替え。',
 mgin='次回番号、タイトル、問い',
 mgcheck='次回テーマ確定後に文言差し替え(現状は仮)。'))

C.append(cut('C51','P5',
 '(ナレーションなし)',
 'エンドカード:「次回:『おすすめ』はなぜ当たるのか(仮)」→「Score: Unsaid Works」→レーベルロゴの順(第1回C55マスター複製)。ピアノがフェードアウト',
 'AEマスター複製', fixed=15, bgm='M3',
 telop=None,
 mgid='AE_END_CARD',
 mgauto='エンドカードマスター複製。次回文言のみ差し替え。',
 mgin='次回文言'))

# ---------- TC計算 ----------
def fmt(sec):
    m = int(sec)//60; s = int(sec)%60
    return f"{m}:{s:02d}"

t = 0.0
rows = []
for c in C:
    nar = c['nar']
    if c['fixed']:
        dur = c['fixed']
    else:
        chars = len(nar.replace('\n','').replace('★',''))
        dur = max(5.0, math.ceil(chars/CPS + c['extra']))
    tc = f"{fmt(t)}-{fmt(t+dur)}"
    rows.append([c['no'], c['no'][0]+'', tc])  # placeholder
    c['_tc'] = tc
    t += dur
total = t
print(f"合計尺: {fmt(total)} ({len(C)}カット)")

# ---------- ワークブック書き込み ----------
wb = openpyxl.load_workbook(SRC)
ws = wb['コンテ']
ws.delete_rows(2, ws.max_row)

r = 2
for c in C:
    vals = [c['no'], c['part'], c['_tc'], c['nar'], c['yomi'], c['visual'], c['source'],
            c['envato'], c['runway'], c['runway_ng'], c['bgm'], c['se'], c['se_q'],
            c['fx'], c['trn'], c['telop'], c['tmotion'], c['tfx'], c['mgid'],
            c['mgauto'], c['mgin'], c['mgtool'], c['mgcheck'], c['mgjson']]
    for col, v in enumerate(vals, start=1):
        ws.cell(row=r, column=col, value=v)
    r += 1

# ---------- 全体設計シート更新 ----------
ws2 = wb['全体設計']
updates = {
 '回': '第2回「怒りは、燃料になる」——あなたの怒りは、誰の収益になっているのか',
 '台本': 'v0.1ドラフト(Fable作成。ナレーション全文は仮。収録前に要確定)',
 '尺': f'約{fmt(total)}(ナレーション約{sum(len(c["nar"].replace(chr(10),"").replace("★","")) for c in C)}字/5.5字・秒)',
 '素材ソース凡例': 'AEマスター複製=第1回で承認済みのAEコンポを複製し文字・素材のみ差し替え/Remotion生成=remotion_mg/のデータMGキットで透過書き出し→AE仕上げ/HTML/UI生成=架空UIをHTMLで組み画面収録/Canva+AI静止画=Canva組版+AI画像のヒーロー素材/Envato=Envato Elements/AI生成=Runway/実撮影=Yuki撮影 ※本回から「自作MG(AE新規構築)」は原則廃止',
 '法務チェック(公開前)': '①実在チャンネル・アカウント・投稿の再現禁止(炎上事例はすべて架空に再構成)②国内事案は固有名なし・三点留保(本人否定/報道側訂正/未確定)を維持 ③SNS内部文書は「報道による」留保形 ④拡散研究・怒り学習研究は留保形+概要欄出典 ⑤ハーストの逸話は「とされる」留保 ⑥レイジベイト例に実在人物を連想させる特徴(口癖・構図・料理名等)を入れない ⑦サムネコピーの本編裏付け確認',
}
for row in ws2.iter_rows(min_row=1):
    key = row[0].value
    if key in updates:
        row[1].value = updates[key]

wb.save(DST)
print(f"saved: {DST}")
for c in C[:5]:
    print(c['no'], c['_tc'], c['source'])
