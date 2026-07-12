#!/usr/bin/env python3
"""コンテブックへ地図MGガイドと素材ソース選択肢を追加する。"""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


WORKBOOK = Path("第1回_コンテ表_v1.xlsx")
SECTION_TITLE = "地図MG（構造表現）"
SECTION_TEXT = "\n".join(
    (
        "用途：政策・規制変更が風景を変えた日付、人口移動、地政学的構造の可視化",
        "主な使用回：#4（高層ビル規制）、#8（人口集中）、#1（バーネイズの1920年代）",
        "素材源：Google Earth Studio（権利確認必須）、国土地理院地図（商用OK、過去空中写真あり）",
        "美学：Johnny Harris風の「元気な地図」ではなく、分解図と同じく低彩度・線描・金茶アクセントのみ",
        "ガイドライン：「地図も分解図の一種」と位置づけ、キャラクターを統一すること",
    )
)
SOURCE_CHOICES = (
    "MGテンプレート", "MG自動生成", "HTML/UI生成", "Envatoテンプレート",
    "Envato実写+AE合成", "AI静止画+2.5D", "画面収録", "固有MG",
    "Envato", "実撮影(推奨)", "AI生成(Runway)", "地図MG", "国土地理院GIS",
)


def copy_cell_style(source, target) -> None:
    """既存の全体設計表と同じ見た目を新しい行へ引き継ぐ。"""
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def main() -> None:
    workbook = load_workbook(WORKBOOK)

    # 全体設計は「項目／内容」の2列表なので、独立した1行セクションとして追加する。
    design = workbook["全体設計"]
    existing_row = next(
        (row for row in range(2, design.max_row + 1) if design.cell(row, 1).value == SECTION_TITLE),
        None,
    )
    target_row = existing_row or design.max_row + 1
    if not existing_row:
        copy_cell_style(design.cell(target_row - 1, 1), design.cell(target_row, 1))
        copy_cell_style(design.cell(target_row - 1, 2), design.cell(target_row, 2))
    design.cell(target_row, 1).value = SECTION_TITLE
    design.cell(target_row, 2).value = SECTION_TEXT
    content_alignment = copy(design.cell(target_row, 2).alignment)
    content_alignment.wrap_text = True
    content_alignment.vertical = "top"
    design.cell(target_row, 2).alignment = content_alignment
    title_alignment = copy(design.cell(target_row, 1).alignment)
    title_alignment.vertical = "top"
    design.cell(target_row, 1).alignment = title_alignment
    design.row_dimensions[target_row].height = 90

    # 現ブックにはG列の既存DataValidationがない。G2:G1000へ設定しておけば、
    # 次回シートを複製した場合や行を追加した場合にもドロップダウンが利用できる。
    conte = workbook["コンテ"]
    for validation in list(conte.data_validations.dataValidation):
        if "G2:G1000" in str(validation.sqref):
            conte.data_validations.dataValidation.remove(validation)
    formula = '"' + ",".join(SOURCE_CHOICES) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.promptTitle = "素材ソース"
    validation.prompt = "素材ソースを選択してください。必要に応じて既存値を直接入力することもできます。"
    validation.errorTitle = "素材ソース"
    # 既存55カットにはEnvato等の自由記述があるため、既存値を妨げない設定にする。
    validation.showErrorMessage = False
    validation.showInputMessage = True
    conte.add_data_validation(validation)
    validation.add("G2:G1000")

    workbook.save(WORKBOOK)
    print(f"更新しました: {WORKBOOK}")


if __name__ == "__main__":
    main()
