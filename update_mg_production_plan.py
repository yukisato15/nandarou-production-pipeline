#!/usr/bin/env python3
"""コンテの「自作MG」を具体的な制作方式へ再分類し、詳細計画を追記する。"""

from __future__ import annotations

from copy import copy
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


WORKBOOK = Path("第1回_コンテ表_v1.xlsx")
PLAN_SHEET = "MG制作計画"

# categoryはG列へ入る管理用の正式値。補助ツールは「完成手段」ではなく、
# Canvaで静止部品を作る、Geminiで文字なし素材を作る、といった役割まで明記する。
PLANS = {
    "C01": {
        "category": "MGテンプレート", "template": "AE_CARD_STACK_THUMBNAIL",
        "automation": "3枚のサムネ画像を読み込み、通知音に合わせて順次スタック。各カードの出現遅延、105→100%スケール、最終微ズームを自動設定。",
        "inputs": "サムネ画像3枚、表示順、出現間隔、背景色、最終ズーム量",
        "support": "Canvaで架空サムネの静止画を一括作成。Claude/Codexは画像リストをAEへ流し込むJSX生成。VHSノイズはEnvato部品。",
        "manual": "サムネ文言の可読性と煽りの強さだけ確認。実在チャンネル・ロゴが混入していないことを確認。",
        "reuse": "第2回以降の冒頭フックに毎回転用可能。最優先で番組用AEマスターへ登録。",
    },
    "C02": {
        "category": "HTML/UI生成", "template": "UI_VIDEO_FEED_HOVER",
        "automation": "架空動画フィードをHTML/CSSで描画し、スクロール停止、ホバー、カーソルのためらい、選択カード暗転をJSONから生成。",
        "inputs": "カード画像、タイトル、再生数、スクロール速度、停止対象、カーソル座標",
        "support": "Claude/CodexでHTML/CSS/JavaScriptを保守。C01のCanva製サムネをカード画像として再利用。",
        "manual": "実在YouTube UIの完全複製にせず、架空サービスとして法務・見た目を確認。",
        "reuse": "SNS・動画フィード系の基本レンダラー。C26、C30へ派生。",
    },
    "C03": {
        "category": "MGテンプレート", "template": "AE_CARD_FOCUS_DROP",
        "automation": "C01のカード部品を再利用し、指定サムネを中央へ落下、周囲カードを暗転、静止後にモノクロ化。",
        "inputs": "主役サムネ画像、落下時間、周囲暗転率、モノクロ化時間",
        "support": "主役画像はCanvaまたは実際の本動画サムネ。微ダストのみEnvato。",
        "manual": "C01からC03までのカード位置が連続して見えるか確認。",
        "reuse": "カード選択・注目演出として全話で再利用可能。",
    },
    "C05": {
        "category": "HTML/UI生成", "template": "UI_ANALYTICS_DASHBOARD",
        "automation": "架空分析ダッシュボード、折れ線、CTR、再生数、カード比較をHTML/SVGで生成し、数値変動を自動アニメーション。",
        "inputs": "指標名、開始値・終了値、グラフ点、サムネ画像、変動速度",
        "support": "Claude/CodexでUIコード生成。本人PC画面を使う場合は画面収録へ切替可能。",
        "manual": "YouTube Studioそのものに見えすぎないこと、個人情報・実数値が露出しないことを確認。",
        "reuse": "分析、A/Bテスト、カウンター画面の共通UIキット。C19、C44と共用。",
    },
    "C06": {
        "category": "Envatoテンプレート", "template": "ENV_SCREEN_WALL_SWAP",
        "automation": "No PluginのTV／Screen Wallテンプレートを番組色へ改造し、各画面素材と切替タイミングだけ差し替える。",
        "inputs": "TV画面素材、スマホ画面素材、切替フレーム、点滅強度",
        "support": "Envatoのscreen wall／video wallを部品採用。Claude/Codexはプレースホルダー一括差替えJSXを作成。",
        "manual": "外部プラグイン不要か確認し、点滅が強すぎないよう調整。素材ライセンスを当該回へ登録。",
        "reuse": "画面群・メディア比較用のEnvato改造マスターとして保存。",
    },
    "C10": {
        "category": "MGテンプレート", "template": "AE_TITLE_EPISODE",
        "automation": "話数、タイトル、副題を差し替え、字間収束、紙質感、フェード、次カットへの粒子化を番組プリセットで適用。",
        "inputs": "話数、タイトル、副題、保持秒数、アウト方式",
        "support": "番組用AEマスターのみ。紙テクスチャはEnvatoライブラリから固定素材を利用。",
        "manual": "長い題名の改行と保持時間のみ確認。",
        "reuse": "全エピソード共通。最優先で作るブランドテンプレート。",
    },
    "C12": {
        "category": "MG自動生成", "template": "DATA_BAR_WORD_GROWTH",
        "automation": "JSONの語句と値から棒グラフShape Layerを生成。棒の伸長、ラベル、数値カウント、順位変化を自動設定。",
        "inputs": "語句配列、各値、表示順、アクセント対象、最大値、尺",
        "support": "PythonでMG JSONを生成し、AE Shape Layer／Trim Pathsで描画。データ案の整理にClaudeを利用。",
        "manual": "値が事実データか演出値かを明記し、誤認を招かないか確認。",
        "reuse": "ランキング・比較・増減説明に汎用利用できるデータMGキット。",
    },
    "C13": {
        "category": "MG自動生成", "template": "DATA_FLOW_THREE_STEP",
        "automation": "3段のノード、矢印、アイコン、数値をJSONから生成し、視線→再生→広告費を順番に開示。",
        "inputs": "ノード名3件、補足値、アイコンID、強調ノード、表示順",
        "support": "Claude/CodexでSVGアイコンとJSXを生成。共通アイコンは番組ライブラリへ保存。",
        "manual": "因果関係を断定しすぎない文言か確認。",
        "reuse": "仕組み・因果・商流説明の基本モジュール。",
    },
    "C15": {
        "category": "MG自動生成", "template": "DATA_RESEARCH_STAT_CARD",
        "automation": "出典、年、件数をJSONから論文風カードへ配置し、大数字のカウントアップとスタンプ着地を自動生成。",
        "inputs": "研究機関、年、件数、単位、注記、出典URL、表示順",
        "support": "AEデータカードマスター＋Python。Canvaは不要。紙テクスチャのみEnvato固定部品。",
        "manual": "数字・単位・出典の正確性、約表記、留保文を必ず確認。",
        "reuse": "研究紹介、統計、年代表示に高頻度で再利用可能。",
    },
    "C16": {
        "category": "MG自動生成", "template": "DATA_LINE_COMPARE_TWO",
        "automation": "2系列の点列・ラベル・色から線グラフを生成。Trim Pathsで描画し、一方が追い抜くタイミングと終点ラベルを自動設定。",
        "inputs": "系列名2件、点配列、色、追い抜き時刻、軸表示有無、注記",
        "support": "AE Shape Layer中心。試作が重い場合は無料CavalryのSpreadsheet連携も比較候補。",
        "manual": "軸を省略する場合の誤解、実測値と概念図の区別を確認。",
        "reuse": "A/B比較、増減、時系列比較の中心モジュール。",
    },
    "C17": {
        "category": "HTML/UI生成", "template": "UI_MESSAGE_BRANCH_NETWORK",
        "automation": "メッセージ文、送信者、枝分かれ数、遅延をJSONからCanvas/SVGへ展開し、吹き出しが連鎖増殖する動画を生成。",
        "inputs": "メッセージ配列、分岐数、世代数、送信間隔、最大表示数",
        "support": "Claude/CodexでCanvasロジックを構築。通知音はEnvato。",
        "manual": "画面密度、読み取れる主メッセージ、点滅速度を確認。",
        "reuse": "拡散、ネットワーク、口コミ表現へ転用可能。初回構築はやや重い。",
    },
    "C18": {
        "category": "固有MG", "template": "CUSTOM_MESSAGE_TO_HUMAN_CIRCUIT",
        "automation": "C17の吹き出し座標を受け取り、AE上の人型パスへ吸着・収束させる補助スクリプトまでは自動化。最終モーフは固有調整。",
        "inputs": "吹き出し座標、人型SVG、収束時間、回路線本数、発光量",
        "support": "Gemini等で文字なしの回路図風人型案を生成し、Claude/CodexでSVGを整理。最終アニメーションはAE。",
        "manual": "人型として認識できる瞬間と、C17からの連続性を手動調整。",
        "reuse": "完全自動化対象外。ただしノード収束エンジンは別テーマでも再利用可能。",
    },
    "C19": {
        "category": "HTML/UI生成", "template": "UI_AB_TEST_DASHBOARD",
        "automation": "C05のUIキットを使い、サムネ2案、CTR、勝敗、変化率をJSONから自動描画。",
        "inputs": "案名、画像2枚、CTR2値、勝者、試行数、表示時間",
        "support": "C05のHTML/SVGコンポーネントを再利用。",
        "manual": "架空データ表記と可読性を確認。",
        "reuse": "A/Bテスト・比較UIとして横展開。",
    },
    "C20": {
        "category": "MGテンプレート", "template": "AE_BLACK_DECLARATION",
        "automation": "黒背景、S4中央文、10Fイン・アウト、保持時間を既存テロップJSONから自動配置。",
        "inputs": "本文、保持時間、フェードF数",
        "support": "既存generate_ae_layers.jsxを利用。新規素材不要。",
        "manual": "黒の純度と無音の間だけ確認。",
        "reuse": "C40と同一。全話の強い宣言カットに利用。",
    },
    "C21": {
        "category": "Envatoテンプレート", "template": "ENV_TIME_REWIND_FILM",
        "automation": "Old Film／Film Burn素材と調整レイヤーをプリセット化し、カラー→白黒→セピアの変化を指定尺へ自動フィット。",
        "inputs": "元映像、変化開始・終了、粒状感、ゲート揺れ量",
        "support": "Envatoのfilm leader、film burn、gate weave。",
        "manual": "素材ライセンス登録、過度な古写真表現にならないか確認。",
        "reuse": "時代遡行・歴史パート導入の共通トランジション。",
    },
    "C25": {
        "category": "MG自動生成", "template": "DATA_REPLACE_ARROW_DIAGRAM",
        "automation": "左語句を打ち消し、右語句へ矢印が伸びるミニマル図解を文字列からShape Layerとして生成。",
        "inputs": "左語句、右語句、打消し方式、矢印方向、アクセント色",
        "support": "Claude/CodexがJSX／SVGを生成。紙質感は番組AEマスター。",
        "manual": "左右の意味関係と改行だけ確認。",
        "reuse": "定義変更、誤解→本質、Before→Afterに使える。",
    },
    "C26": {
        "category": "HTML/UI生成", "template": "UI_AD_FEED_SCROLL",
        "automation": "架空ブランド広告カードをデータから生成し、スマホフィードへ挿入。高速スクロールと指定カードでの停止を自動化。",
        "inputs": "ブランド名、画像、見出し、本文、CTA、配色、カード順",
        "support": "Canva Bulk Createで広告静止カード案を作成可能。量産時はClaude/Codex製HTMLテンプレートを優先。Gemini画像は文字なし商品背景のみ。",
        "manual": "実在ブランド類似、AI画像の権利・不自然さ、広告文の可読性を確認。",
        "reuse": "C02フィードレンダラーの広告モード。",
    },
    "C30": {
        "category": "HTML/UI生成", "template": "UI_LIST_TO_SOCIAL_FEED",
        "automation": "3項テキストを順次スタンプ表示し、同じ文言をSNSカードへ引き継いでタイムラインへマッチカット。",
        "inputs": "3項目、カード投稿者、投稿文、表示順、スタンプ間隔",
        "support": "前半はAEリストマスター、後半はC02のHTMLフィード。Claude/Codexが接続データを生成。",
        "manual": "3項目の改行、紙面→SNSのマッチ位置を確認。",
        "reuse": "箇条書きから実例UIへ移る場面に利用。",
    },
    "C31": {
        "category": "Envato実写+AE合成", "template": "AE_MEDIA_EVOLUTION_MONTAGE",
        "automation": "輪転機・ラジオ・TV・スマホの4スロットへ実写を入れ、時代別LUT、尺、ホイップパンをマスターコンポで一括適用。",
        "inputs": "実写4本、各開始点、時代ラベル、色調プリセット",
        "support": "Envato実写を検索・取得。AEは4枠モンタージュマスター。",
        "manual": "各実写の年代感、日本人が出る場合の検索条件、ライセンスを確認。",
        "reuse": "技術・制度・メディアの変遷モンタージュへ転用。",
    },
    "C33": {
        "category": "Envato実写+AE合成", "template": "AE_PHONE_CROWD_FOCUS",
        "automation": "暗所スマホ実写へ画面差替えプリコンポを適用し、1台だけフォーカス・発光・カメラ寄りを自動設定。",
        "inputs": "実写、注目端末の座標、差替え画面、寄り量、発光量",
        "support": "日本人を含むEnvato実写。画面内容はHTML/UI生成。",
        "manual": "注目端末のトラッキングと画面四隅だけ確認。",
        "reuse": "多数の中の1人／1台を強調する合成マスター。",
    },
    "C34": {
        "category": "HTML/UI生成", "template": "UI_DEVICE_GRID_ADS",
        "automation": "端末数、列数、奥行き、広告カード配列をJSONから生成し、各画面を順次点灯。AE側でグリッドを3D配置。",
        "inputs": "端末数、行列、広告データ、点灯順、カメラ移動、強調数字",
        "support": "広告カードはCanvaでラフ量産可能。最終文字はHTMLで正確に描画。Claude/Codexが生成器を保守。",
        "manual": "情報量と処理負荷、広告の重複、画面外はみ出しを確認。",
        "reuse": "多数端末・パーソナライズ表現の中核UIモジュール。",
    },
    "C35": {
        "category": "HTML/UI生成", "template": "UI_DEVICE_COMPARE_TWO",
        "automation": "C34を2列比較モードで再利用し、同じ候補に対する2種類の広告、配色、見出しを左右へ生成。",
        "inputs": "共通候補名、左広告、右広告、左右配色、強調語",
        "support": "Claude/Codex製UIコンポーネント。背景画像が必要ならGeminiで文字なし素材を作成。",
        "manual": "左右の差が瞬時に伝わるか、色だけに依存していないか確認。",
        "reuse": "不安／怒り、A/B、属性別配信の比較に利用。",
    },
    "C36": {
        "category": "Envatoテンプレート", "template": "ENV_NEWSPAPER_STACK_GENERIC",
        "automation": "No Pluginの新聞／紙面スタックへ架空見出しと年を流し込み、重なり、印刷ノイズ、カメラ引きを共通化。",
        "inputs": "架空見出し配列、日付、年、紙面画像、重なり順",
        "support": "Envato新聞テンプレートを部品化。Canvaで架空紙面の静止レイアウトを作り、最終文字はAEで置換。",
        "manual": "実在媒体に見えないこと、誤読される見出しでないこと、ライセンスを確認。",
        "reuse": "報道・資料・歴史見出しの共通テンプレート。",
    },
    "C38": {
        "category": "Envato実写+AE合成", "template": "AE_CITY_PHONE_LIGHTS",
        "automation": "夜の街・スマホ実写へ低彩度LUT、光ブルーム、点描マスク、S2配置をプリセット適用。",
        "inputs": "実写、発光ポイント、LUT、ブルーム量、カメラ寄り",
        "support": "日本人が映るEnvato実写を使用。新規MGは最小限。",
        "manual": "人物属性、スマホ光の位置、次の値札変形への接続を確認。",
        "reuse": "都市と端末光を扱う実写合成プリセット。",
    },
    "C40": {
        "category": "MGテンプレート", "template": "AE_BLACK_DECLARATION",
        "automation": "C20と同じコンポを複製せず同一マスターから生成し、同じ位置・文字・モーションを保証。",
        "inputs": "C20参照ID、保持時間",
        "support": "既存JSXのみ。",
        "manual": "C20との完全な同一性を確認。",
        "reuse": "回収演出のための参照機能もMG Factoryへ実装。",
    },
    "C42": {
        "category": "MGテンプレート", "template": "AE_INTERSTITIAL_SIMPLE",
        "automation": "白／墨背景、短文、フェード、次画面の横スライドをパラメータだけで生成。",
        "inputs": "本文、背景色、文字色、保持秒、アウト方向",
        "support": "番組用AEマスター。素材不要。",
        "manual": "箸休めとして長すぎないか確認。",
        "reuse": "章間、転換、短い問いに使用。",
    },
    "C43": {
        "category": "Envato実写+AE合成", "template": "AE_MENU_OVERLAY_GENERIC",
        "automation": "ロゴなし店舗実写へ架空メニューUIを合成し、商品名・価格・蛍光灯ちらつきをデータ差替え。",
        "inputs": "実写、メニュー文、価格、画面座標、ちらつき量",
        "support": "日本人スタッフ／客のEnvato実写。メニュー静止案はCanva、最終文字はAEまたはHTML。",
        "manual": "実在チェーン類似、看板・ロゴ、合成パース、ライセンスを確認。",
        "reuse": "店舗メニュー・券売機・価格表示の合成キット。",
    },
    "C44": {
        "category": "HTML/UI生成", "template": "UI_SOCIAL_SCORE_COUNTER",
        "automation": "フォロワー数と架空スコアをJSONから描画し、カウントアップ、停止、評価バーを生成。",
        "inputs": "開始値、終了値、単位、プロフィール画像、スコア、停止時刻",
        "support": "C05ダッシュボードのカウンター部品を再利用。Claude/Codexで生成。",
        "manual": "実在アプリに見えすぎないこと、数値の意味が誤解されないか確認。",
        "reuse": "フォロワー、評価、価格、スコア全般に利用。",
    },
    "C45": {
        "category": "MGテンプレート", "template": "AE_DARK_STATEMENT",
        "automation": "暗色背景、S2中央テロップ、静止保持、値札タグ出現を番組マスターから生成。",
        "inputs": "本文、保持秒、値札表示有無、値札出現時刻",
        "support": "既存テロップJSX＋値札部品。",
        "manual": "言葉だけに集中できる余白と保持時間を確認。",
        "reuse": "黒画面宣言のS2版。",
    },
    "C46": {
        "category": "Envato実写+AE合成", "template": "AE_PRICE_TAG_MULTIPLY",
        "automation": "実写上の指定点へ値札プリコンポを複製し、出現順、価格、スケール、密度をJSONから設定。最後に一括消去。",
        "inputs": "実写、座標配列、値札文言、出現間隔、最大数、消去時刻",
        "support": "日本人の日常Envato実写＋AE Repeater／JSX。値札デザインはCanvaで案出し可能だが正本はAE。",
        "manual": "顔や重要被写体を隠さない位置、増殖テンポ、ライセンスを確認。",
        "reuse": "値札、ラベル、注釈、追跡マーカーの増殖表現に利用。",
    },
    "C48": {
        "category": "Envato実写+AE合成", "template": "AE_LIST_ON_AMBIENT_PLATE",
        "automation": "光の移ろい実写へ3項目リストマスターを重ね、項目ごとの出現、既出項の減光、背景ディゾルブを自動設定。",
        "inputs": "背景実写、項目配列、出現時刻、既出濃度、位置",
        "support": "Envatoの静かな室内光実写＋AEリストテンプレート。",
        "manual": "背景が主張しすぎないか、3項目の改行と読み時間を確認。",
        "reuse": "まとめ、手順、論点整理の標準カット。",
    },
    "C52": {
        "category": "MGテンプレート", "template": "AE_BRAND_LOGO_ASSEMBLY",
        "automation": "番組ロゴ、分解図線画、紙質感、線描アニメ、文字フェードをブランドマスター化し、話数や副題のみ差替え可能にする。",
        "inputs": "ロゴSVG、話数、副題、線画バリエーション、保持秒",
        "support": "Gemini等は文字なし分解図モチーフの案出しのみ。Claude/CodexでSVG整理と描画JSXを作り、最終デザインはAEに固定。",
        "manual": "初回だけデザインを丁寧に決定。以後は変更せずブランド一貫性を優先。",
        "reuse": "全話共通の最重要資産。初期工数をかける価値が最大。",
    },
    "C53": {
        "category": "画面収録", "template": "CAPTURE_NOTE_SCROLL",
        "automation": "公開用noteまたはローカルHTMLを自動スクロールし、規定速度で画面収録。必要ならブラウザ枠だけAEで付与。",
        "inputs": "記事URL／HTML、開始位置、終了位置、スクロール速度、見せる出典箇所",
        "support": "本人PC画面収録。Claude/Codexは公開前のローカルHTMLモックと自動スクロールスクリプトを作成可能。",
        "manual": "個人情報、通知、ブラウザタブ、未公開情報、出典表示を確認。",
        "reuse": "自動化より実画面収録が早い例外。ブラウザ枠プリセットのみ再利用。",
    },
    "C54": {
        "category": "MGテンプレート", "template": "AE_NEXT_EPISODE_TITLE",
        "automation": "次回番号、題名、問いを差し替え、黒背景、粒子ループ、フェード、保持を自動適用。",
        "inputs": "次回番号、タイトル、問い、保持秒、粒子強度",
        "support": "粒子はEnvatoの固定ループをライブラリ化。Gemini画像生成は不要。",
        "manual": "次回題名の確定、改行、予告尺を確認。",
        "reuse": "全話共通の次回予告マスター。",
    },
    "C55": {
        "category": "MGテンプレート", "template": "AE_END_CARD_CREDITS",
        "automation": "次回表示、Score表記、レーベルロゴをJSONから順次配置し、3秒間隔クロスフェードと黒落ちを生成。",
        "inputs": "次回文、Score名義、ロゴ画像、各保持秒、終了秒",
        "support": "番組用AEマスター。外部素材不要。",
        "manual": "クレジット表記、音楽終了位置、YouTube終了画面との干渉を確認。",
        "reuse": "全話共通。最優先で作るエンドカード資産。",
    },
    # 「自作MG」の文字はないが、同じ管理体系に整理すると迷いが減る2カット。
    "C14": {
        "category": "画面収録", "template": "CAPTURE_EDIT_TIMELINE_DEMO",
        "automation": "編集前後の操作を本人PCで収録し、画面枠・ズーム・注釈・カーソル強調だけ共通AEプリセットを適用。",
        "inputs": "収録動画、Before／After時刻、強調座標、注釈文",
        "support": "本人PC画面収録。Claude/Codexはデモ用素材と注釈リスト作成を補助。",
        "manual": "個人情報、ファイル名、通知、他案件素材を完全に隠す。",
        "reuse": "操作そのものは手動、見せ方だけテンプレート化。",
    },
    "C27": {
        "category": "AI静止画+2.5D", "template": "AI_PROPAGANDA_BOOK_PARALLAX",
        "automation": "文字なし古書・紙・型押し素材を生成し、前景／表紙／背景へ分離。AEカメラマスターで寄り、光なぞり、視差を適用。",
        "inputs": "生成静止画、表紙文字、前景マスク、カメラ移動、光方向",
        "support": "Gemini等で文字なし静止画を生成。正確なPROPAGANDA文字はAEで追加。Claude/Codexでマスク・レイヤー整理を補助。",
        "manual": "歴史表現、記号・ロゴ混入、文字誤生成、画像権利、時代考証を確認。",
        "reuse": "古書・資料・象徴物の2.5D見せ方を共通化。",
    },
}


SOURCE_CHOICES = (
    "MGテンプレート", "MG自動生成", "HTML/UI生成", "Envatoテンプレート",
    "Envato実写+AE合成", "AI静止画+2.5D", "画面収録", "固有MG",
    "Envato", "実撮影(推奨)", "AI生成(Runway)", "地図MG", "国土地理院GIS",
)
DETAIL_HEADERS = (
    "MGテンプレートID", "MG自動化内容", "MG必要入力", "MG補助ツール・素材",
    "MG手動確認・注意", "MGパラメータJSON",
)

# MG Factoryで実装済みのカットに流し込む実データ。
# 色は番組パレット、時刻は各カット内の相対フレームで持ち、最終配置時に
# 23.976fpsのTCへ変換する。事実値でない数値はconceptual=trueを必須にする。
IMPLEMENTED_PARAMS = {
    "C01": {
        "fictional": True,
        "card_size": [1040, 360], "pop_frames": 8,
        "final_zoom": 104, "final_zoom_start": 180,
        "cards": [
            {"eyebrow": "速報", "headline": "日本、終わる", "accent": "#B8352E", "x": 0.37, "y": 0.34, "rotation": -3, "delay": 0},
            {"eyebrow": "知らないと危険", "headline": "年金、消える", "accent": "#C9A063", "x": 0.63, "y": 0.50, "rotation": 2, "delay": 18},
            {"eyebrow": "損する前に", "headline": "知らないと、損する", "accent": "#EAE6DF", "x": 0.49, "y": 0.67, "rotation": -1, "delay": 36},
        ],
    },
    "C03": {
        "fictional": True,
        "kicker": "THIS VIDEO", "headline": "同じ手口で作っています",
        "subline": "『なんだろう』解体 #01",
        "drop_frames": 18, "settle_frames": 8,
        "surrounding_opacity": 16, "monochrome_frame": 78,
    },
    "C10": {
        "episode": "#01", "series": "『なんだろう』解体",
        "title": "ぜんぶ、売り物になっていく",
        "subtitle": "あなたの「つい」は、誰のお金になっているのか",
        "intro_frames": 18, "outro_frames": 18,
    },
    "C15": {
        "data_kind": "factual", "source_verified": False,
        "kicker": "MIT RESEARCH", "year": "2018", "value": "約12万件",
        "numeric_value": 126000, "count_to": 12, "count_prefix": "約", "count_suffix": "万件",
        "caption": "ニュース記事の拡散を分析",
        "footnote": "MIT研究／約126,000件（コンテ記載値）",
        "sequence_frames": [0, 18, 38],
    },
    "C16": {
        "title": "拡散の広がり方", "conceptual": True,
        "note": "概念図（値は演出用）",
        "series": [
            {"label": "本当", "color": "#EAE6DF", "values": [0.08, 0.18, 0.29, 0.43, 0.56, 0.67]},
            {"label": "ウソ", "color": "#B8352E", "values": [0.05, 0.22, 0.48, 0.70, 0.87, 0.96]},
        ],
        "draw_frames": 82,
    },
    "C12": {
        "title": "不安ワードの『伸び』", "conceptual": True,
        "note": "概念図（値は演出用）", "draw_frames": 62, "stagger_frames": 12,
        "bars": [
            {"label": "経済が終わる", "value": 96, "accent": True},
            {"label": "年金が消える", "value": 84, "accent": False},
            {"label": "知らないと損", "value": 68, "accent": False},
        ],
    },
    "C13": {
        "title": "注意がお金に変わるまで", "conceptual": True,
        "nodes": [
            {"icon": "◉", "label": "視線", "sub": "ATTENTION"},
            {"icon": "▶", "label": "再生", "sub": "VIEW"},
            {"icon": "¥", "label": "広告費", "sub": "REVENUE"},
        ],
        "reveal_frames": [8, 38, 68],
        "statement": "売られているのは、あなたの『注意』です。",
        "statement_frame": 104,
    },
    "C20": {
        "text": "友達は、お金で買えます。", "reuse_key": "FRIENDS_FOR_SALE",
        "fade_frames": 10, "fade_out_frame": 178, "tracking": 150,
    },
    "C25": {
        "kicker": "売っているもの", "left": "商品そのもの", "right": "なれる自分",
        "strike_frame": 28, "arrow_frame": 48, "reveal_frame": 70,
        "statement": "『それを買うと、どんな自分になれるか』を売る。",
    },
    "C26": {
        "service": "DAILY FEED", "scroll_pixels": 980, "scroll_frames": 260,
        "stop_card": 2,
        "cards": [
            {"brand": "LUMEN", "is_fictional": True, "headline": "もっと軽やかな毎日へ", "body": "暮らしの選択を、シンプルに。", "accent": "#C9A063"},
            {"brand": "MIRAI+", "is_fictional": True, "headline": "あなた向けに選びました", "body": "いま気になる情報をまとめてチェック。", "accent": "#B8352E"},
            {"brand": "COMMON", "is_fictional": True, "headline": "同じ商品、違う見せ方", "body": "広告表現の違いを示す架空ブランドです。", "accent": "#6E7770"},
        ],
    },
    "C40": {
        "text": "友達は、お金で買えます。", "reuse_key": "FRIENDS_FOR_SALE",
        "fade_frames": 10, "fade_out_frame": 178, "tracking": 150,
    },
    "C42": {
        "text": "もうひとつ。", "secondary": "笑顔にも、値段がついたことがあります。",
        "background": "#EAE6DF", "foreground": "#1A1A1A",
        "accent": "#C9A063", "fade_frames": 12, "secondary_frame": 54,
        "out_direction": "left",
    },
    "C45": {
        "text": "人間に、市場価値。", "caption": "よく考えると、すごい言葉です。",
        "fade_frames": 12, "caption_frame": 58, "show_price_tag": True,
        "price_tag_frame": 94,
    },
    "C46": {
        "background_mode": "replace_media",
        "background_label": "ここへEnvato日常実写を差し替え",
        "tag_size": [250, 96], "pop_frames": 6,
        "tags": [
            {"label": "通勤", "price": "¥?", "x": 0.16, "y": 0.25, "delay": 0},
            {"label": "時間", "price": "¥?", "x": 0.38, "y": 0.31, "delay": 8},
            {"label": "団らん", "price": "¥?", "x": 0.63, "y": 0.23, "delay": 16},
            {"label": "安心", "price": "¥?", "x": 0.79, "y": 0.40, "delay": 24},
            {"label": "注意", "price": "¥?", "x": 0.22, "y": 0.64, "delay": 32},
            {"label": "感情", "price": "¥?", "x": 0.48, "y": 0.70, "delay": 40},
            {"label": "関係", "price": "¥?", "x": 0.76, "y": 0.68, "delay": 48},
        ],
        "avoid_regions": [],
    },
    "C48": {
        "background_mode": "replace_media", "background_label": "静かな室内光の実写へ差し替え",
        "position": "left", "reveal_frames": [12, 122, 244], "dim_opacity": 32,
        "items": [
            {"head": "ひとつ。", "body": "本当は売り物じゃないものまで、\nどんどん売り物になっています。"},
            {"head": "ふたつ。", "body": "『つい反応してしまう』心を、\n仕組みはよく知って使ってきます。"},
            {"head": "みっつ。", "body": "仕組みは、知れば、距離が取れます。"},
        ],
    },
    "C52": {
        "series": "『なんだろう』解体", "episode": "#01",
        "subtitle": "構造から、ひとつずつ。", "line_count": 8,
        "draw_frames": 70, "logo_frame": 54, "subtitle_frame": 82,
    },
    "C54": {
        "kicker": "NEXT EPISODE", "next_episode": "#02",
        "title": "あなたの怒りは、\n誰の収益になっているのか",
        "question": "SNSを見たあとの、あのモヤモヤ。",
        "fade_frames": 18, "particle_count": 14,
    },
    "C55": {
        "hold_frames": 72, "crossfade_frames": 18, "end_fade_frames": 24,
        "cards": [
            {"kind": "next", "text": "次回：あなたの怒りは誰の収益か"},
            {"kind": "credit", "text": "Score: Unsaid Works"},
            {"kind": "logo", "text": "Unsaid Works"},
        ],
    },
}


def copy_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def main() -> None:
    workbook = load_workbook(WORKBOOK)
    conte = workbook["コンテ"]
    # 再実行時も「旧素材ソース」を失わないよう、既存計画シートから退避する。
    saved_original_sources: dict[str, str] = {}
    if PLAN_SHEET in workbook.sheetnames:
        old_plan = workbook[PLAN_SHEET]
        for row in range(2, old_plan.max_row + 1):
            cut = str(old_plan.cell(row, 1).value or "").strip()
            if cut:
                saved_original_sources[cut] = str(old_plan.cell(row, 4).value or "")
    headers = {str(cell.value).strip(): cell.column for cell in conte[1] if cell.value is not None}
    no_col = headers["No"]
    source_col = headers["素材ソース"]
    part_col = headers["パート"]
    tc_col = headers["TC目安"]
    visual_col = headers["映像内容(画)"]

    # S〜Xを詳細管理欄として追加。ヘッダー名で扱う既存Pythonツールには影響しない。
    detail_start = 19
    for offset, header in enumerate(DETAIL_HEADERS):
        column = detail_start + offset
        conte.cell(1, column).value = header
        copy_style(conte.cell(1, 18), conte.cell(1, column))
        conte.column_dimensions[conte.cell(1, column).column_letter].width = [25, 55, 40, 55, 50, 60][offset]

    original_sources: dict[str, str] = {}
    row_by_cut: dict[str, int] = {}
    for row in range(2, conte.max_row + 1):
        cut = str(conte.cell(row, no_col).value or "").strip()
        row_by_cut[cut] = row
        if cut not in PLANS:
            continue
        plan = PLANS[cut]
        original_sources[cut] = saved_original_sources.get(cut, str(conte.cell(row, source_col).value or ""))
        conte.cell(row, source_col).value = plan["category"]
        params = IMPLEMENTED_PARAMS.get(cut, {})
        details = (
            plan["template"], plan["automation"], plan["inputs"], plan["support"], plan["manual"],
            json.dumps(params, ensure_ascii=False, separators=(",", ":")) if params else "",
        )
        for offset, value in enumerate(details):
            cell = conte.cell(row, detail_start + offset)
            cell.value = value
            copy_style(conte.cell(row, 18), cell)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        conte.row_dimensions[row].height = max(conte.row_dimensions[row].height or 15, 90)

    missing = sorted(set(PLANS) - set(row_by_cut))
    if missing:
        raise ValueError("対象カットが見つかりません: " + ", ".join(missing))

    # G列ドロップダウンを新分類へ更新。既存の自由記述を妨げない設定は維持する。
    for validation in list(conte.data_validations.dataValidation):
        if "G2:G1000" in str(validation.sqref):
            conte.data_validations.dataValidation.remove(validation)
    validation = DataValidation(type="list", formula1='"' + ",".join(SOURCE_CHOICES) + '"', allow_blank=True)
    validation.showErrorMessage = False
    validation.showInputMessage = True
    validation.promptTitle = "素材ソース／制作方式"
    validation.prompt = "MGは制作方式で分類してください。必要なら自由入力も可能です。"
    conte.add_data_validation(validation)
    validation.add("G2:G1000")

    # 一覧シートは毎回作り直し、コンテ側の最新内容と同期する。
    if PLAN_SHEET in workbook.sheetnames:
        del workbook[PLAN_SHEET]
    plan_sheet = workbook.create_sheet(PLAN_SHEET, 2)
    plan_headers = (
        "Cut", "パート", "TC", "旧素材ソース", "新しい制作方式", "テンプレートID",
        "MG Factory実装状況", "現在の映像内容", "自動化する内容", "必要入力データ", "補助ツール・素材の使い所",
        "人が確認すること", "将来の再利用方針",
    )
    for column, header in enumerate(plan_headers, 1):
        cell = plan_sheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A1A1A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ordered_cuts = sorted(PLANS, key=lambda value: int(value[1:]))
    for output_row, cut in enumerate(ordered_cuts, 2):
        source_row = row_by_cut[cut]
        plan = PLANS[cut]
        values = (
            cut,
            conte.cell(source_row, part_col).value,
            conte.cell(source_row, tc_col).value,
            original_sources.get(cut, "（既に更新済み）"),
            plan["category"],
            plan["template"],
            "実装済み v0.2" if cut in IMPLEMENTED_PARAMS else "未実装",
            conte.cell(source_row, visual_col).value,
            plan["automation"],
            plan["inputs"],
            plan["support"],
            plan["manual"],
            plan["reuse"],
        )
        for column, value in enumerate(values, 1):
            cell = plan_sheet.cell(output_row, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if output_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F1ED")
            if column == 7:
                cell.font = Font(bold=True, color="1A1A1A")
                cell.fill = PatternFill("solid", fgColor="D9E8D2" if cut in IMPLEMENTED_PARAMS else "E7E3DC")
        plan_sheet.row_dimensions[output_row].height = 105
    widths = [9, 9, 13, 22, 22, 34, 18, 55, 65, 50, 65, 55, 50]
    for column, width in enumerate(widths, 1):
        plan_sheet.column_dimensions[plan_sheet.cell(1, column).column_letter].width = width
    plan_sheet.freeze_panes = "A2"
    plan_sheet.auto_filter.ref = f"A1:M{plan_sheet.max_row}"
    plan_sheet.sheet_view.showGridLines = False

    workbook.save(WORKBOOK)
    print(f"更新しました: コンテ詳細 {len(PLANS)}カット / {PLAN_SHEET} {len(PLANS)}行")


if __name__ == "__main__":
    main()
