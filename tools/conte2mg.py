#!/usr/bin/env python3
"""コンテxlsxのMG制作情報を、再利用可能なMG専用JSONへ変換する。"""

from __future__ import annotations

import argparse
from collections import Counter
import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from mg_factory.template_registry import IMPLEMENTED_TEMPLATES, validate_cut


REPO_ROOT = Path(__file__).resolve().parent.parent


HEADERS = {
    "cut": "No",
    "part": "パート",
    "tc": "TC目安",
    "visual": "映像内容(画)",
    "category": "素材ソース",
    "envato": "Envato検索ワード(映像)",
    "ai_prompt": "Runway生成プロンプト",
    "effect": "エフェクト(+Envato検索)",
    "transition": "次カットへのトランジション(+Envato検索)",
    "motion": "テロップモーション",
    "template": "MGテンプレートID",
    "automation": "MG自動化内容",
    "inputs": "MG必要入力",
    "support": "MG補助ツール・素材",
    "manual": "MG手動確認・注意",
    "params": "MGパラメータJSON",
}
MG_CATEGORIES = {
    "MGテンプレート", "MG自動生成", "HTML/UI生成", "Envatoテンプレート",
    "Envato実写+AE合成", "AI静止画+2.5D", "画面収録", "固有MG",
}
CUT_RANGE_RE = re.compile(r"^([A-Za-z]+)(\d+)(?:-([A-Za-z]+)(\d+))?$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="コンテxlsxからMG Factory用JSONを生成")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, default=REPO_ROOT / "out/ep01/ep01_mg.json")
    parser.add_argument("--sheet", default="コンテ")
    parser.add_argument("--episode", default="EP01")
    parser.add_argument("--fps", type=float, default=23.976)
    parser.add_argument("--cuts", help="対象範囲（例: C10-C16、単独ならC26）")
    parser.add_argument("--include-planned", action="store_true", help="MG対象外でもテンプレートIDがある行を含める")
    return parser.parse_args()


def cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def tc_seconds(value: str) -> tuple[float, float]:
    match = re.fullmatch(r"\s*(\d+):(\d{1,2}(?:\.\d+)?)\s*-\s*(\d+):(\d{1,2}(?:\.\d+)?)\s*", value)
    if not match:
        raise ValueError(f"TC形式が不正: {value!r}")
    sm, ss, em, es = match.groups()
    if float(ss) >= 60 or float(es) >= 60:
        raise ValueError(f"TC秒は60未満: {value!r}")
    return int(sm) * 60 + float(ss), int(em) * 60 + float(es)


def format_time(seconds: float) -> str:
    millis = round(seconds * 1000)
    hours, millis = divmod(millis, 3_600_000)
    minutes, millis = divmod(millis, 60_000)
    secs, millis = divmod(millis, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"


def cut_filter(specification: str | None):
    if not specification:
        return lambda _: True
    match = CUT_RANGE_RE.fullmatch(specification.strip())
    if not match:
        raise ValueError(f"--cuts形式が不正: {specification!r}")
    start_prefix, start_number, end_prefix, end_number = match.groups()
    end_prefix = end_prefix or start_prefix
    end_number = end_number or start_number
    if start_prefix.upper() != end_prefix.upper():
        raise ValueError("--cutsは同じ接頭辞で指定してください")
    start, end = int(start_number), int(end_number)
    if start > end:
        raise ValueError("--cutsの開始番号は終了番号以下にしてください")

    def accepted(value: str) -> bool:
        found = re.fullmatch(r"([A-Za-z]+)(\d+)", value)
        return bool(found and found.group(1).upper() == start_prefix.upper() and start <= int(found.group(2)) <= end)

    return accepted


def parse_params(cut: str, raw: str, warnings: list[str]) -> dict[str, Any]:
    if not raw:
        return {}
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        warnings.append(f"{cut}: MGパラメータJSON不正 ({exc.msg})")
        return {}
    if not isinstance(value, dict):
        warnings.append(f"{cut}: MGパラメータJSONはobjectにしてください")
        return {}
    return value


def load_mg(args: argparse.Namespace) -> tuple[dict[str, Any], list[str]]:
    workbook = load_workbook(args.input, data_only=True, read_only=True)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"シートがありません: {args.sheet}")
    sheet = workbook[args.sheet]
    header = {cell_text(cell.value): cell.column for cell in next(sheet.iter_rows(min_row=1, max_row=1)) if cell_text(cell.value)}
    missing = [value for value in HEADERS.values() if value not in header]
    if missing:
        raise ValueError("必須ヘッダーがありません: " + ", ".join(missing))
    if args.fps <= 0:
        raise ValueError("--fpsは正数にしてください")

    accepts = cut_filter(args.cuts)
    warnings: list[str] = []
    cuts: list[dict[str, Any]] = []
    get = lambda row, key: cell_text(sheet.cell(row, header[HEADERS[key]]).value)
    for row in range(2, sheet.max_row + 1):
        cut_no = get(row, "cut")
        if not cut_no or not accepts(cut_no):
            continue
        category = get(row, "category")
        template = get(row, "template")
        if category not in MG_CATEGORIES and not (args.include_planned and template):
            continue
        if category not in MG_CATEGORIES:
            warnings.append(f"{cut_no}: 未知のMG制作方式 {category!r}")
        if not template:
            warnings.append(f"{cut_no}: MGテンプレートIDが空です")
            continue
        try:
            start, end = tc_seconds(get(row, "tc"))
        except ValueError as exc:
            warnings.append(f"{cut_no}: {exc}")
            continue
        if start >= end:
            warnings.append(f"{cut_no}: TCが逆転または尺0")
            continue
        in_frame = round(start * args.fps)
        out_frame = round(end * args.fps)
        cut = {
            "cut": cut_no,
            "part": get(row, "part"),
            "tc_in": format_time(start),
            "tc_out": format_time(end),
            "in_seconds": start,
            "out_seconds": end,
            "in_frame": in_frame,
            "out_frame": out_frame,
            # 独立した尺roundではなくTCフレーム差にして、親配置と1Fずれないようにする。
            "duration_frames": out_frame - in_frame,
            "category": category,
            "template_id": template,
            "visual": get(row, "visual"),
            "automation": get(row, "automation"),
            "required_inputs": get(row, "inputs"),
            "support": get(row, "support"),
            "manual_check": get(row, "manual"),
            "source_search": get(row, "envato"),
            "ai_prompt": get(row, "ai_prompt"),
            "effect_notes": get(row, "effect"),
            "transition_notes": get(row, "transition"),
            "motion_notes": get(row, "motion"),
            "params": parse_params(cut_no, get(row, "params"), warnings),
        }
        if template in IMPLEMENTED_TEMPLATES:
            warnings.extend(validate_cut(cut))
        cuts.append(cut)

    payload = {
        "schema_version": "1.0",
        "project": {
            "episode": args.episode,
            "title": "『なんだろう』解体",
            "width": 1920,
            "height": 1080,
            "fps": args.fps,
            "fps_numerator": 24000 if abs(args.fps - 23.976) < 0.01 else None,
            "fps_denominator": 1001 if abs(args.fps - 23.976) < 0.01 else None,
            "duration_seconds": max((cut["out_seconds"] for cut in cuts), default=0),
        },
        "palette": {
            "ink": "#1A1A1A",
            "off_white": "#EAE6DF",
            "gold": "#C9A063",
            "vermilion": "#B8352E",
        },
        "cuts": cuts,
    }
    return payload, warnings


def main() -> int:
    args = parse_args()
    if not args.input.is_absolute():
        args.input = REPO_ROOT / args.input
    if not args.output.is_absolute():
        args.output = REPO_ROOT / args.output
    try:
        payload, warnings = load_mg(args)
    except (OSError, ValueError) as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        return 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    categories = Counter(cut["category"] for cut in payload["cuts"])
    implemented = [cut["cut"] for cut in payload["cuts"] if cut["template_id"] in IMPLEMENTED_TEMPLATES]
    print("=== MG JSON生成レポート ===")
    print(f"MGカット数: {len(payload['cuts'])}")
    print("制作方式: " + ", ".join(f"{key}={value}" for key, value in sorted(categories.items())))
    print(f"実装済み対象: {', '.join(implemented) if implemented else 'なし'}")
    print(f"警告数: {len(warnings)}")
    for warning in warnings:
        print(f"警告: {warning}")
    print(f"出力: {args.output}")
    return 0 if not warnings else 0


if __name__ == "__main__":
    raise SystemExit(main())
