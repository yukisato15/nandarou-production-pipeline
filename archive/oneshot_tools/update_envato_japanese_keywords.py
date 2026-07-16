#!/usr/bin/env python3
"""人物が出るEnvato映像検索語を、日本人を優先する語へ更新する。"""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path("第1回_コンテ表_v1.xlsx")

# Envatoは英語検索を前提とし、人・手・家族・群衆が映り得る検索には
# Japaneseを明記する。建物や純粋な物撮りだけの検索語は変更しない。
KEYWORDS = {
    "C04": "\n".join((
        "Japanese broadcast camera operator newsroom",
        "Japanese video editor hands editing suite",
        "Japanese hands news script paper desk",
    )),
    "C09": "\n".join((
        "Japanese researcher desk notes coffee hands",
        "Japanese person writing notebook closeup",
    )),
    "C11": "\n".join((
        "Japanese person doomscrolling dark room phone glow",
        "Japanese hand scrolling smartphone closeup night",
    )),
    "C14": "Japanese people town street evening neutral",
    "C31": "\n".join((
        "vintage radio closeup",
        "retro television static",
        "Japanese person holding smartphone glowing",
    )),
    "C32": "\n".join((
        "Japanese family watching television vintage retro",
        "Japanese people reading newspaper 1970s",
    )),
    "C33": "Japanese people using many smartphones dark room glowing",
    "C38": "\n".join((
        "Tokyo night Japanese people using smartphones",
        "Japanese crowd looking at phones night",
    )),
    "C39": "\n".join((
        "Japanese supermarket checkout register",
        "price tags closeup Japanese retail store",
        "Japanese people waiting in line supermarket",
    )),
    "C41": "\n".join((
        "Japanese wedding reception guests clapping",
        "Japanese wedding banquet hall guests",
    )),
    "C43": "\n".join((
        "Japanese fast food counter staff no logo generic",
        "Japanese burger restaurant customers menu board generic",
    )),
    "C46": "\n".join((
        "Japanese commuters walking morning Tokyo",
        "Japanese family dinner table daily life",
        "Japanese person park bench everyday life",
    )),
    "C47": "\n".join((
        "morning light Japanese home window curtain",
        "quiet Japanese street soft morning light",
        "steam coffee cup Japanese home morning",
    )),
    "C49": "\n".join((
        "Japanese magician card trick hands closeup",
        "Japanese hands card sleight of hand slow motion",
    )),
    "C51": "\n".join((
        "Tokyo evening street Japanese people walking warm",
        "Japanese people city dusk daily life distant",
    )),
}


def main() -> None:
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["コンテ"]
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
    no_column = headers["No"]
    keyword_column = headers["Envato検索ワード(映像)"]
    updated = []
    for row in range(2, sheet.max_row + 1):
        cut = str(sheet.cell(row, no_column).value or "").strip()
        if cut in KEYWORDS:
            sheet.cell(row, keyword_column).value = KEYWORDS[cut]
            updated.append(cut)
    missing = sorted(set(KEYWORDS) - set(updated))
    if missing:
        raise ValueError("対象カットが見つかりません: " + ", ".join(missing))
    workbook.save(WORKBOOK)
    print(f"更新しました: {len(updated)}カット ({', '.join(updated)})")


if __name__ == "__main__":
    main()
